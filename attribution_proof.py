"""Pure evidence preparation and XLSX export for the Hydrogen attribution proof."""
from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

import pandas as pd

from cost_allocation.resource_allocation import (
    UPSTREAM_SUPPLY_VIEW,
    ResourceAllocationResult,
    with_conversion_column_names,
)
from cost_allocation.resource_cascade import (
    CASCADE_ROLE_FINAL,
    CASCADE_ROLE_PASSED_ON,
    CASCADE_ROLE_UNALLOCATED,
)

HYDROGEN_POOL = "Hydrogen"
_TOLERANCE = 1e-6
_PROCESS_TAXONOMY_COLUMNS = [
    "process_group", "process_category", "process_subsector",
    "process_sector", "process_ghgi", "process_dukes", "process_useful_output", "process_set",
]
_CONNECTION_ROLE_COLUMNS = ["connection_role", "connection_role_reason"]
HYDROGEN_SANKEY_AGGREGATIONS = {
    "Individual processes": None,
    "Technology group": "process_group",
    "Process category": "process_category",
    "Subsector": "process_subsector",
    "Sector": "process_sector",
    "GHGI": "process_ghgi",
    "DUKES": "process_dukes",
    "Useful output": "process_useful_output",
    "Process set": "process_set",
}
DEFAULT_HYDROGEN_SOURCE_AGGREGATION = "Technology group"
DEFAULT_HYDROGEN_USE_AGGREGATION = "DUKES"
FINANCIAL_SIGN_COSTS = "Costs"
FINANCIAL_SIGN_CREDITS = "Credits & revenues"
FINANCIAL_SIGN_ZERO = "Zero / non-material"
STAGE_SOURCE = "Supply-chain cost source"
STAGE_PASSED_IN = "Conversion supply"
STAGE_USE = "Use-category attribution"
STAGE_PASSED_ON = "Conversion demand"
STAGE_UNALLOCATED = "Unallocated"
HYDROGEN_ATTRIBUTION_WARNING = (
    "This is an attribution of existing upstream model costs, not a direct cost booked on a "
    "receiving technology, a physical ownership trace, or an extra system-cost total. Do not "
    "add it to overlapping direct-cost or Flow-cost views."
)


@dataclass(frozen=True)
class HydrogenProof:
    """Complete, presentation-ready evidence for one run and model period."""

    run_id: str
    period: int
    sources: pd.DataFrame
    source_costs: pd.DataFrame
    users: pd.DataFrame
    excluded_flows: pd.DataFrame
    sector_shares: pd.DataFrame
    financial_links: pd.DataFrame
    physical_links: pd.DataFrame
    vintage_detail: pd.DataFrame
    checks: pd.DataFrame
    diagnostics: pd.DataFrame
    # DUKES use share for each share-basis year (the build year for investment).
    basis_shares: pd.DataFrame = field(default_factory=lambda: pd.DataFrame(
        columns=["share_basis_year", "process_dukes", "dukes_share"]
    ))


def _empty(columns: list[str]) -> pd.DataFrame:
    return pd.DataFrame(columns=columns)


def _numeric(frame: pd.DataFrame, column: str) -> pd.Series:
    return pd.to_numeric(frame.get(column, pd.Series(index=frame.index, dtype=float)), errors="coerce").fillna(0.0)


def _text(frame: pd.DataFrame, column: str, default: str = "") -> pd.Series:
    return frame.get(column, pd.Series(default, index=frame.index)).fillna(default).astype(str).str.strip()


def _vdt_supply_processes(bundle) -> frozenset[str]:
    """Return the selected run's normalised VDT-derived supply boundary."""
    scope = getattr(bundle, "energy_scope", None)
    return frozenset(
        str(process).strip().upper()
        for process in getattr(scope, "supply_processes", ())
        if str(process).strip()
    )


def hydrogen_sankey_aggregation_field(selection: str) -> str | None:
    """Return the authoritative label field for a displayed Sankey level."""
    try:
        return HYDROGEN_SANKEY_AGGREGATIONS[selection]
    except KeyError as exc:
        options = ", ".join(HYDROGEN_SANKEY_AGGREGATIONS)
        raise ValueError(f"Unsupported Hydrogen Sankey aggregation: {selection}. Choose one of: {options}.") from exc


def _display_category(frame: pd.DataFrame, field: str | None) -> pd.Series:
    """Return an external-facing label without changing the underlying process evidence."""
    if field is None:
        return pd.Series(
            [_source_label(process, name) for process, name in zip(
                _text(frame, "process"), _text(frame, "process_name")
            )],
            index=frame.index,
        ).replace("", "Unlabelled process")
    return _text(frame, field).replace("", "Unlabelled")


def _node_id(stage: str, value: object) -> str:
    """Build a stage-qualified node identity independent of its display label."""
    return f"{stage}|{str(value or '').strip().upper() or 'UNLABELLED'}"


FINANCIAL_SIGN_NET = "Net"


def net_financial_links(links: pd.DataFrame) -> pd.DataFrame:
    """Net cost and credit links per displayed link; draw net credits reversed.

    ``value`` keeps its sign. A reversed link's non-pool end gets its own node
    identity so it does not merge with a cost node of the same name.
    """
    columns = [
        "source_id", "source", "target_id", "target", "value", "unit", "stage",
        "financial_sign", "hover_detail", "gross_cost_m", "gross_credit_m",
    ]
    if links.empty:
        return pd.DataFrame(columns=columns)
    rows = links.copy()
    rows["value"] = pd.to_numeric(rows["value"], errors="coerce").fillna(0.0)
    rows["gross_cost_m"] = rows["value"].where(rows["value"].gt(0), 0.0)
    rows["gross_credit_m"] = rows["value"].where(rows["value"].lt(0), 0.0)
    keys = ["stage", "source_id", "source", "target_id", "target"]
    net = rows.groupby(keys, as_index=False, sort=False).agg(
        value=("value", "sum"),
        gross_cost_m=("gross_cost_m", "sum"),
        gross_credit_m=("gross_credit_m", "sum"),
        unit=("unit", "first"),
    )
    net = net[net["value"].abs().gt(_TOLERANCE)].copy()
    pool_id = _node_id("financial_pool", HYDROGEN_POOL)
    reverse = net["value"].lt(0)
    flipped = net.loc[reverse].copy()
    for end in ("source_id", "target_id"):
        moved = flipped[end].ne(pool_id)
        flipped.loc[moved, end] = flipped.loc[moved, end] + "|net_credit"
    net.loc[reverse, "source_id"] = flipped["target_id"]
    net.loc[reverse, "target_id"] = flipped["source_id"]
    net.loc[reverse, "source"] = flipped["target"]
    net.loc[reverse, "target"] = flipped["source"]
    net["financial_sign"] = FINANCIAL_SIGN_NET
    net["hover_detail"] = [
        f"{'Net credit (drawn reversed)' if value < 0 else 'Net cost'}; "
        f"gross cost {cost:,.3f}, gross credit {credit:,.3f}"
        for value, cost, credit in zip(net["value"], net["gross_cost_m"], net["gross_credit_m"])
    ]
    return net.loc[:, columns].reset_index(drop=True)


def _contributor_detail(frame: pd.DataFrame) -> str:
    """Summarise process evidence for an aggregated external-facing node hover."""
    contributors = frame[["process", "process_name"]].drop_duplicates().copy()
    contributor_labels = [
        _source_label(row.process, row.process_name)
        for row in contributors.itertuples(index=False)
    ]
    preview = "; ".join(contributor_labels[:8])
    remainder = len(contributor_labels) - 8
    if remainder > 0:
        preview += f"; +{remainder} more"
    return f"Contributing processes ({len(contributor_labels)}): {preview or 'None'}"


def available_hydrogen_periods(result: ResourceAllocationResult) -> list[int]:
    """Return periods with a Hydrogen source, user, or allocated amount."""
    candidates: list[pd.Series] = []
    for frame in (result.cost_assignments, result.quantity_lineage, result.allocations):
        if not frame.empty and {"period", "resource_pool"}.issubset(frame.columns):
            rows = frame[_text(frame, "resource_pool").eq(HYDROGEN_POOL)]
            candidates.append(pd.to_numeric(rows["period"], errors="coerce"))
    if not candidates:
        return []
    values = pd.concat(candidates, ignore_index=True).dropna().astype(int).unique().tolist()
    return sorted(values)


def preferred_hydrogen_period(result: ResourceAllocationResult) -> int | None:
    """Prefer 2030 when it has evidence, otherwise use the earliest evidence."""
    periods = available_hydrogen_periods(result)
    return 2030 if 2030 in periods else (periods[0] if periods else None)


def _source_outputs(
    result: ResourceAllocationResult,
    bundle,
    period: int,
    source_costs: pd.DataFrame,
) -> pd.DataFrame:
    """Identify net Hydrogen source injections per (region, process).

    A process is a physical source only for its positive net injection:
    ``max(Hydrogen OUT − Hydrogen IN, 0)`` at the same region, process and
    period.  Distribution processes that merely pass Hydrogen through have
    zero net injection and do not appear as sources.
    """
    lineage = result.quantity_lineage.copy()
    columns = [
        "process", "process_name", *_PROCESS_TAXONOMY_COLUMNS, *_CONNECTION_ROLE_COLUMNS,
        "commodity", "commodity_name",
        "observed_output_twh", "observed_input_twh", "net_source_twh", "physical_role",
        "source_cost_scope", "source_cost_scope_reason", "assigned_source_cost_m",
    ]
    if lineage.empty:
        return _empty(columns)
    hydrogen = lineage[
        _text(lineage, "resource_pool").eq(HYDROGEN_POOL)
        & pd.to_numeric(lineage.get("period"), errors="coerce").eq(period)
    ].copy()
    if hydrogen.empty:
        return _empty(columns)
    hydrogen["process"] = _text(hydrogen, "process")
    hydrogen["region"] = _text(hydrogen, "region")
    hydrogen["direction"] = _text(hydrogen, "direction").str.upper()
    hydrogen["quantity_twh"] = _numeric(hydrogen, "quantity_twh").abs()
    supply = _vdt_supply_processes(bundle)

    # Aggregate Hydrogen IN and OUT per (region, process) for supply processes.
    out_rows = hydrogen[
        hydrogen["direction"].eq("OUT") & hydrogen["process"].isin(supply)
    ].copy()
    in_rows = hydrogen[
        hydrogen["direction"].eq("IN") & hydrogen["process"].isin(supply)
    ].copy()
    if out_rows.empty:
        return _empty(columns)
    out_totals = out_rows.groupby(
        ["region", "process"], as_index=False, dropna=False
    )["quantity_twh"].sum().rename(columns={"quantity_twh": "observed_output_twh"})
    in_totals = in_rows.groupby(
        ["region", "process"], as_index=False, dropna=False
    )["quantity_twh"].sum().rename(columns={"quantity_twh": "observed_input_twh"})
    balance = out_totals.merge(in_totals, on=["region", "process"], how="left")
    balance["observed_input_twh"] = balance["observed_input_twh"].fillna(0.0)
    balance["net_source_twh"] = (balance["observed_output_twh"] - balance["observed_input_twh"]).clip(lower=0.0)
    balance["physical_role"] = balance["net_source_twh"].apply(
        lambda value: "ultimate_source" if value > _TOLERANCE else "throughput_or_delivery"
    )
    # Only processes with positive net injection are physical sources.
    sources = balance[balance["net_source_twh"] > _TOLERANCE].copy()
    if sources.empty:
        return _empty(columns)

    # Attach process metadata from the first OUT row per (region, process).
    metadata_columns = [
        "region", "process", "process_name",
        *_PROCESS_TAXONOMY_COLUMNS, *_CONNECTION_ROLE_COLUMNS,
        "commodity", "commodity_name",
    ]
    meta = out_rows.copy()
    for column in metadata_columns:
        if column not in meta.columns:
            meta[column] = ""
        meta[column] = _text(meta, column)
    meta = meta.drop_duplicates(["region", "process"], keep="first")[metadata_columns]
    sources = sources.merge(meta, on=["region", "process"], how="left")
    sources["source_cost_scope"] = "VDT-derived Hydrogen supply"
    sources["source_cost_scope_reason"] = (
        "This net Hydrogen source injection belongs to the selected run's VDT-derived "
        "supply boundary. Hydrogen input at the same process is deducted first."
    )
    assigned = _empty(["process", "assigned_source_cost_m"])
    if not source_costs.empty:
        assigned = source_costs.assign(
            process=_text(source_costs, "process"),
            assigned_source_cost_m=_numeric(source_costs, "assigned_cost_m"),
        ).groupby("process", as_index=False)["assigned_source_cost_m"].sum()
    sources = sources.merge(assigned, on="process", how="left").fillna({"assigned_source_cost_m": 0.0})
    for column in columns:
        if column not in sources.columns:
            sources[column] = "" if column not in {"observed_output_twh", "observed_input_twh", "net_source_twh", "assigned_source_cost_m"} else 0.0
    return sources.loc[:, columns].sort_values(
        ["net_source_twh", "process"], ascending=[False, True]
    ).reset_index(drop=True)

def _source_costs(result: ResourceAllocationResult, bundle, period: int) -> pd.DataFrame:
    columns = [
        "source_row_id", "process", "process_name", *_PROCESS_TAXONOMY_COLUMNS,
        "attribute", "vintage", "vintage_basis", "share_basis_year", "share_basis",
        "cost_value_m", "assignment_share", "assigned_cost_m", "assignment_status", "assignment_reason",
        "cost_scope", "cost_scope_reason",
    ]
    source = result.cost_assignments.copy()
    if source.empty:
        return _empty(columns)
    source = source[
        _text(source, "view").eq(UPSTREAM_SUPPLY_VIEW)
        & _text(source, "resource_pool").eq(HYDROGEN_POOL)
        & pd.to_numeric(source.get("period"), errors="coerce").eq(period)
    ].copy()
    if source.empty:
        return _empty(columns)
    source["process"] = _text(source, "process")
    source = source[source["process"].isin(_vdt_supply_processes(bundle))].copy()
    if source.empty:
        return _empty(columns)
    source["cost_scope"] = "VDT supply-chain Hydrogen cost"
    source["cost_scope_reason"] = (
        "This assigned upstream Hydrogen cost belongs to the selected run's "
        "VDT-derived supply boundary. Contributors include production, network, "
        "distribution, storage and other route activity."
    )
    for column in columns:
        if column not in source.columns:
            source[column] = pd.NA if column in {"vintage", "share_basis_year"} else ""
    return source.loc[:, columns].sort_values(["process", "attribute", "source_row_id"], na_position="last").reset_index(drop=True)


def _hydrogen_inputs(
    lineage: pd.DataFrame, period: int, end_use: set[str],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return the included and excluded observed Hydrogen inputs for one period."""
    hydrogen = lineage[
        _text(lineage, "resource_pool").eq(HYDROGEN_POOL)
        & _text(lineage, "direction").str.upper().eq("IN")
        & pd.to_numeric(lineage.get("period"), errors="coerce").eq(period)
    ].copy()
    for column in (
        "process", "process_name", "sector", "commodity", "commodity_name",
        *_PROCESS_TAXONOMY_COLUMNS, *_CONNECTION_ROLE_COLUMNS,
    ):
        hydrogen[column] = _text(hydrogen, column, "Unlabelled" if column in {"process_name", "sector"} else "")
    hydrogen["observed_input_twh"] = _numeric(hydrogen, "quantity_twh")
    hydrogen["structural_end_use"] = hydrogen["process"].str.upper().isin(end_use)
    included = hydrogen[hydrogen.get("included", pd.Series(False, index=hydrogen.index)).fillna(False).astype(bool)].copy()
    excluded = hydrogen[~hydrogen.index.isin(included.index)].copy()
    return included, excluded


def _basis_years(source_costs: pd.DataFrame) -> pd.Series:
    """Share-basis year per source cost row; missing where the row is unallocated."""
    return pd.to_numeric(
        source_costs.get("share_basis_year", pd.Series(pd.NA, index=source_costs.index)),
        errors="coerce",
    )


def _hydrogen_ledger(result: ResourceAllocationResult, period: int) -> pd.DataFrame:
    """Engine ledger rows entering or leaving the Hydrogen pool, with process labels.

    The ledger is the allocation itself: final use, cost passed on by
    converters and cost with no route to final use, each row already split by
    its own share-basis year's conversion chain.
    """
    columns = [
        "sign", "resource_pool", "region", "process", "process_name",
        *_PROCESS_TAXONOMY_COLUMNS, "sector", "role", "passed_to_pool", "allocated_cost_m",
    ]
    ledger = result.process_allocations
    if ledger is None or ledger.empty:
        return _empty(columns)
    rows = ledger[
        _text(ledger, "view").eq(UPSTREAM_SUPPLY_VIEW)
        & pd.to_numeric(ledger.get("period"), errors="coerce").eq(period)
    ].copy()
    for column in ("sign", "resource_pool", "region", "process", "sector", "role", "passed_to_pool"):
        rows[column] = _text(rows, column)
    rows = rows[
        rows["resource_pool"].eq(HYDROGEN_POOL)
        | (rows["role"].eq(CASCADE_ROLE_PASSED_ON) & rows["passed_to_pool"].eq(HYDROGEN_POOL))
    ]
    rows["allocated_cost_m"] = _numeric(rows, "allocated_cost_m")
    rows = rows[rows["allocated_cost_m"].abs().gt(_TOLERANCE)]
    lineage = result.quantity_lineage
    label_columns = ["process_name", *_PROCESS_TAXONOMY_COLUMNS]
    if lineage.empty:
        labels = _empty(["region", "process", *label_columns])
    else:
        labels = lineage.copy()
        for column in ("region", "process", *label_columns):
            labels[column] = _text(labels, column)
        # Prefer this period's labels; a build-year user may only appear earlier.
        labels["_other_period"] = ~pd.to_numeric(labels.get("period"), errors="coerce").eq(period)
        labels = labels.sort_values("_other_period", kind="stable").drop_duplicates(
            ["region", "process"]
        ).loc[:, ["region", "process", *label_columns]]
    rows = rows.drop(columns=[c for c in label_columns if c in rows.columns]).merge(
        labels, on=["region", "process"], how="left"
    )
    for column in label_columns:
        rows[column] = rows[column].fillna("")
    return rows.loc[:, columns].reset_index(drop=True)


def _users(
    result: ResourceAllocationResult,
    bundle,
    period: int,
    source_costs: pd.DataFrame,
    ledger: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Return reporting-year users, exclusions, DUKES attribution and basis-year shares.

    ``allocated_cost_m`` and ``passed_on_cost_m`` come from the engine ledger:
    Hydrogen used by a process that turns it into another pool (a hydrogen
    turbine, an ammonia plant) is passed on, and only final use stays with a
    DUKES classification.  ``sector_twh`` and ``allocation_share`` stay the
    reporting year's use evidence.  ``basis_shares`` replays the first step of
    the split — a row's cost times a DUKES class's share of Hydrogen use in the
    row's share-basis year — for the zoom-in.
    """
    user_columns = [
        "process", "process_name", *_PROCESS_TAXONOMY_COLUMNS, *_CONNECTION_ROLE_COLUMNS,
        "sector", "commodity", "commodity_name", "observed_input_twh",
        "structural_end_use", "sector_twh", "pool_twh", "allocation_share",
        "allocated_cost_m", "passed_on_cost_m",
    ]
    share_columns = [
        "process_dukes", "sector_twh", "pool_twh", "allocation_share",
        "allocated_cost_m", "passed_on_cost_m", "cost_share",
    ]
    lineage = result.quantity_lineage.copy()
    if lineage.empty:
        return (
            _empty(user_columns),
            _empty(["process", "commodity", "quantity_status", "quantity_twh"]),
            _empty(share_columns),
            _empty(["share_basis_year", "process_dukes", "dukes_share"]),
        )
    end_use = {str(value).strip().upper() for value in getattr(getattr(bundle, "energy_scope", None), "end_use_processes", ())}
    included, excluded = _hydrogen_inputs(lineage, period, end_use)
    basis_costs = _numeric(source_costs, "assigned_cost_m").groupby(_basis_years(source_costs)).sum()
    own_hydrogen = ledger["resource_pool"].eq(HYDROGEN_POOL)
    passed_in = float(ledger.loc[
        ~own_hydrogen & ledger["role"].eq(CASCADE_ROLE_PASSED_ON), "allocated_cost_m"
    ].sum())
    pool_cost = float(_numeric(source_costs, "assigned_cost_m").sum()) + passed_in
    reporting = (
        included.groupby("process_dukes", dropna=False)["observed_input_twh"].sum()
        if not included.empty else pd.Series(dtype=float)
    )
    allocated = ledger[own_hydrogen & ledger["role"].eq(CASCADE_ROLE_FINAL)].groupby(
        "process_dukes"
    )["allocated_cost_m"].sum()
    passed_on = ledger[own_hydrogen & ledger["role"].eq(CASCADE_ROLE_PASSED_ON)].groupby(
        "process_dukes"
    )["allocated_cost_m"].sum()
    basis_share_rows: list[dict[str, object]] = []
    for year in basis_costs.index:
        year_users = included if int(year) == period else _hydrogen_inputs(lineage, int(year), end_use)[0]
        by_dukes = year_users.groupby("process_dukes", dropna=False)["observed_input_twh"].sum()
        total = float(by_dukes.sum())
        if abs(total) <= _TOLERANCE:
            continue
        for dukes, quantity in by_dukes.items():
            basis_share_rows.append({
                "share_basis_year": int(year), "process_dukes": dukes,
                "dukes_share": quantity / total,
            })
    basis_shares = pd.DataFrame(
        basis_share_rows, columns=["share_basis_year", "process_dukes", "dukes_share"]
    )
    dukes_values = sorted(
        set(reporting.index).union(allocated.index).union(passed_on.index), key=str
    )
    if not dukes_values:
        dukes_shares = _empty(share_columns)
    else:
        pool_twh = float(reporting.sum())
        dukes_shares = pd.DataFrame({"process_dukes": dukes_values})
        dukes_shares["sector_twh"] = dukes_shares["process_dukes"].map(reporting).fillna(0.0)
        dukes_shares["pool_twh"] = pool_twh
        dukes_shares["allocation_share"] = (
            dukes_shares["sector_twh"] / pool_twh if abs(pool_twh) > _TOLERANCE else 0.0
        )
        dukes_shares["allocated_cost_m"] = dukes_shares["process_dukes"].map(allocated).fillna(0.0)
        dukes_shares["passed_on_cost_m"] = dukes_shares["process_dukes"].map(passed_on).fillna(0.0)
        # The class's share of all cost entering the pool, kept or passed on.
        dukes_shares["cost_share"] = (
            (dukes_shares["allocated_cost_m"] + dukes_shares["passed_on_cost_m"]) / pool_cost
            if abs(pool_cost) > _TOLERANCE else float("nan")
        )
    if included.empty:
        users = _empty(user_columns)
    else:
        users = included.merge(dukes_shares, on="process_dukes", how="left")
        for column in ("sector_twh", "pool_twh", "allocation_share", "allocated_cost_m", "passed_on_cost_m"):
            users[column] = _numeric(users, column)
        users = users.loc[:, user_columns].sort_values(["sector", "process", "commodity"]).reset_index(drop=True)
    excluded_columns = [
        "process", "process_name", *_PROCESS_TAXONOMY_COLUMNS, *_CONNECTION_ROLE_COLUMNS,
        "sector", "commodity", "commodity_name", "quantity_status",
        "observed_input_twh", "structural_end_use",
    ]
    if excluded.empty:
        excluded_display = _empty(excluded_columns)
    else:
        excluded_display = excluded.copy()
        excluded_display["quantity_status"] = _text(excluded_display, "quantity_status")
        excluded_display = excluded_display.loc[:, excluded_columns].sort_values(["quantity_status", "process"])
    sector_shares = dukes_shares.sort_values("process_dukes").reset_index(drop=True)
    return users, excluded_display, sector_shares, basis_shares


def _source_label(process: object, process_name: object) -> str:
    """Keep a short, unique Sankey label while retaining the model code."""
    code = str(process or "").strip()
    name = str(process_name or "").strip()
    return f"{name} ({code})" if name and name != code else code


def _taxonomy_detail(row: object) -> str:
    """Return compact, explicit source taxonomy for graph hover evidence."""
    values = []
    for column, label in (
        ("process_group", "Technology group"),
        ("process_category", "Process category"),
        ("process_subsector", "Subsector"),
        ("process_sector", "NZS / Costs sector"),
        ("process_dukes", "DUKES"),
        ("process_useful_output", "Useful output"),
    ):
        value = str(getattr(row, column, "") or "").strip()
        if value:
            values.append(f"{label}: {value}")
    return "<br>".join(values) or "No process taxonomy metadata is available."


def _connection_role_detail(row: object) -> str:
    role = str(getattr(row, "connection_role", "") or "").strip()
    reason = str(getattr(row, "connection_role_reason", "") or "").strip()
    if not role:
        return "Connection role: not available in this evidence frame."
    return f"Connection role: {role}<br>{reason}" if reason else f"Connection role: {role}"


def _source_cost_detail(source_rows: pd.DataFrame) -> str:
    """Summarise source attributes for an aggregated financial Sankey link."""
    details: list[str] = []
    for attribute, rows in source_rows.groupby("attribute", sort=True, dropna=False):
        amount = float(_numeric(rows, "assigned_cost_m").sum())
        text = f"{attribute}: £{amount:,.3f}M/year"
        if str(attribute) in {"Cost_Inv", "Cost_Fom"} and "vintage" in rows.columns:
            vintages = sorted({
                str(int(value)) if float(value).is_integer() else str(value)
                for value in pd.to_numeric(rows["vintage"], errors="coerce").dropna()
            })
            if vintages:
                text += f" (reported vintage: {', '.join(vintages)})"
        details.append(text)
    return "<br>".join(details) or "No eligible source-cost rows."


def _links(
    sources: pd.DataFrame,
    source_costs: pd.DataFrame,
    users: pd.DataFrame,
    ledger: pd.DataFrame,
    *,
    source_aggregation: str,
    use_aggregation: str,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Build a presentation-only aggregation of exact Hydrogen proof evidence."""
    columns = [
        "source_id", "source", "target_id", "target", "value", "unit", "stage",
        "financial_sign", "hover_detail",
    ]
    source_field = hydrogen_sankey_aggregation_field(source_aggregation)
    use_field = hydrogen_sankey_aggregation_field(use_aggregation)
    physical: list[dict[str, Any]] = []
    financial: list[dict[str, Any]] = []
    physical_pool_id = _node_id("physical_pool", HYDROGEN_POOL)
    financial_pool_id = _node_id("financial_pool", HYDROGEN_POOL)

    if not sources.empty:
        source_rows = sources.copy()
        source_rows["_display"] = _display_category(source_rows, source_field)
        for label, rows in source_rows.groupby("_display", sort=True, dropna=False):
            value = float(_numeric(rows, "net_source_twh").sum())
            gross = float(_numeric(rows, "observed_output_twh").sum())
            deducted = float(_numeric(rows, "observed_input_twh").sum())
            physical.append({
                "source_id": _node_id("physical_source", label), "source": label,
                "target_id": physical_pool_id, "target": HYDROGEN_POOL,
                "value": value, "unit": "TWh/year", "stage": "Net Hydrogen source",
                "financial_sign": "",
                "hover_detail": (
                    f"Net Hydrogen source injection: {value:,.3f} TWh/year<br>"
                    f"Gross output: {gross:,.3f}, same-process input deducted: {deducted:,.3f}<br>"
                    f"Displayed by: {source_aggregation}<br>{_contributor_detail(rows)}"
                ),
            })

    if not users.empty:
        user_rows = users.copy()
        user_rows["_display"] = _display_category(user_rows, use_field)
        for label, rows in user_rows.groupby("_display", sort=True, dropna=False):
            quantity = float(_numeric(rows, "observed_input_twh").sum())
            physical.append({
                "source_id": physical_pool_id, "source": HYDROGEN_POOL,
                "target_id": _node_id("physical_use", label), "target": label,
                "value": quantity, "unit": "TWh/year", "stage": "Included Hydrogen end use",
                "financial_sign": "",
                "hover_detail": (
                    f"Observed included Hydrogen input: {quantity:,.3f} TWh/year<br>"
                    f"Displayed by: {use_aggregation}<br>{_contributor_detail(rows)}"
                ),
            })

    # The engine ledger already splits each cost row by its own share-basis
    # year's conversion chain; the proof only groups it for display.
    if not ledger.empty:
        own = ledger["resource_pool"].eq(HYDROGEN_POOL)
        stages = (
            (STAGE_USE, own & ledger["role"].eq(CASCADE_ROLE_FINAL)),
            (STAGE_PASSED_ON, own & ledger["role"].eq(CASCADE_ROLE_PASSED_ON)),
            (STAGE_UNALLOCATED, own & ledger["role"].eq(CASCADE_ROLE_UNALLOCATED)),
            (STAGE_PASSED_IN, ~own),
        )
        for stage, mask in stages:
            stage_rows = ledger[mask].copy()
            if stage_rows.empty:
                continue
            if stage in {STAGE_USE, STAGE_PASSED_ON}:
                stage_rows["_display"] = _display_category(stage_rows, use_field)
            elif stage == STAGE_PASSED_IN:
                stage_rows["_display"] = _display_category(stage_rows, source_field)
            else:
                stage_rows["_display"] = STAGE_UNALLOCATED
            for (label, financial_sign), rows in stage_rows.groupby(
                ["_display", "sign"], sort=True, dropna=False
            ):
                value = float(_numeric(rows, "allocated_cost_m").sum())
                if abs(value) <= _TOLERANCE:
                    continue
                # One node per label on each side, whatever the stage.
                node = _node_id(
                    "financial_source" if stage == STAGE_PASSED_IN else "financial_use", label,
                )
                if stage == STAGE_PASSED_IN:
                    link = {"source_id": node, "source": label,
                            "target_id": financial_pool_id, "target": HYDROGEN_POOL}
                    detail = (
                        f"Supply to Hydrogen: £{value:,.3f}M/year<br>"
                        f"Displayed by: {source_aggregation}"
                    )
                else:
                    link = {"source_id": financial_pool_id, "source": HYDROGEN_POOL,
                            "target_id": node, "target": label}
                    detail = {
                        STAGE_USE: f"Demand on Hydrogen: £{value:,.3f}M/year",
                        STAGE_PASSED_ON: f"Demand on Hydrogen: £{value:,.3f}M/year",
                        STAGE_UNALLOCATED: (
                            f"Hydrogen {financial_sign.lower()} with no route to final use: "
                            f"£{value:,.3f}M/year"
                        ),
                    }[stage]
                if stage in {STAGE_USE, STAGE_PASSED_ON}:
                    detail += f"<br>Displayed by: {use_aggregation}"
                if _text(rows, "process").ne("").any():
                    detail += f"<br>{_contributor_detail(rows[_text(rows, 'process').ne('')])}"
                financial.append({
                    **link, "value": value, "unit": "£M/year", "stage": stage,
                    "financial_sign": financial_sign, "hover_detail": detail,
                })

    if not source_costs.empty:
        cost_rows = source_costs.copy()
        cost_rows["_display"] = _display_category(cost_rows, source_field)
        amounts = _numeric(cost_rows, "assigned_cost_m")
        cost_rows["_financial_sign"] = FINANCIAL_SIGN_ZERO
        cost_rows.loc[amounts.gt(_TOLERANCE), "_financial_sign"] = FINANCIAL_SIGN_COSTS
        cost_rows.loc[amounts.lt(-_TOLERANCE), "_financial_sign"] = FINANCIAL_SIGN_CREDITS
        for (label, financial_sign), rows in cost_rows.groupby(
            ["_display", "_financial_sign"], sort=True, dropna=False
        ):
            total = float(_numeric(rows, "assigned_cost_m").sum())
            if abs(total) <= _TOLERANCE:
                continue
            financial.append({
                "source_id": _node_id("financial_source", label), "source": label,
                "target_id": financial_pool_id, "target": HYDROGEN_POOL,
                "value": total, "unit": "£M/year", "stage": STAGE_SOURCE,
                "financial_sign": financial_sign,
                "hover_detail": (
                    f"VDT supply-chain Hydrogen cost: £{total:,.3f}M/year<br>"
                    f"Contributors may be production, network, distribution, storage or other route activity.<br>"
                    f"Displayed by: {source_aggregation}<br>{_contributor_detail(rows)}<br><br>"
                    f"{_source_cost_detail(rows)}"
                ),
            })
    return pd.DataFrame(physical, columns=columns), pd.DataFrame(financial, columns=columns)


def _checks(
    source_costs: pd.DataFrame,
    sector_shares: pd.DataFrame,
    sources: pd.DataFrame,
    users: pd.DataFrame,
    financial_links: pd.DataFrame,
    result: ResourceAllocationResult,
    bundle,
    period: int,
    source_aggregation: str,
) -> pd.DataFrame:
    checks: list[dict[str, Any]] = []

    def stage_total(links: pd.DataFrame, *stages: str) -> float:
        if links.empty:
            return 0.0
        return float(_numeric(links[_text(links, "stage").isin(stages)], "value").sum())

    entering = stage_total(financial_links, STAGE_SOURCE, STAGE_PASSED_IN)
    leaving = stage_total(financial_links, STAGE_USE, STAGE_PASSED_ON, STAGE_UNALLOCATED)
    denominator = float(_numeric(sector_shares, "pool_twh").max()) if not sector_shares.empty else 0.0
    user_twh = float(_numeric(users, "observed_input_twh").sum())
    share_total = float(_numeric(sector_shares, "allocation_share").sum())
    net_source_twh = float(_numeric(sources, "net_source_twh").sum())
    gross_output_twh = float(_numeric(sources, "observed_output_twh").sum())
    source_input_twh = float(_numeric(sources, "observed_input_twh").sum())
    def add(name: str, actual: float, expected: float, explanation: str) -> None:
        difference = actual - expected
        checks.append({
            "Check": name, "Status": "Pass" if abs(difference) <= _TOLERANCE else "Warning",
            "Actual": actual, "Expected": expected, "Difference": difference, "Detail": explanation,
        })
    add(
        "Hydrogen source cost balances to final DUKES classifications", leaving, entering,
        "Supply to Hydrogen equals demand on it (DUKES final use and conversion) "
        "plus unallocated cost.",
    )
    for financial_sign in (FINANCIAL_SIGN_COSTS, FINANCIAL_SIGN_CREDITS):
        sign_links = financial_links[
            _text(financial_links, "financial_sign").eq(financial_sign)
        ] if not financial_links.empty else financial_links
        source_total = stage_total(sign_links, STAGE_SOURCE, STAGE_PASSED_IN)
        use_total = stage_total(sign_links, STAGE_USE, STAGE_PASSED_ON, STAGE_UNALLOCATED)
        if abs(source_total) > _TOLERANCE or abs(use_total) > _TOLERANCE:
            add(
                f"Hydrogen {financial_sign.lower()} balance to final DUKES classifications",
                use_total,
                source_total,
                f"Signed {financial_sign.lower()} supplied to Hydrogen equal "
                "the demand on it plus unallocated.",
            )
    add("Hydrogen input lineage matches pool denominator", user_twh, denominator, "Included Hydrogen user inputs match the denominator used for DUKES shares.")
    if denominator > _TOLERANCE:
        add("Hydrogen DUKES shares sum to one", share_total, 1.0, "All included DUKES shares should exhaust a positive Hydrogen denominator.")
    else:
        checks.append({"Check": "Hydrogen DUKES shares sum to one", "Status": "Warning", "Actual": share_total, "Expected": 1.0, "Difference": share_total - 1.0, "Detail": "Hydrogen has no positive allocation denominator."})
    checks.append({"Check": "Net source injection versus included end use", "Status": "Information", "Actual": net_source_twh, "Expected": user_twh, "Difference": net_source_twh - user_twh, "Detail": f"Net source injection (gross {gross_output_twh:,.3f} minus same-process input {source_input_twh:,.3f}) and included end use are independent measures; storage, losses, exports, conversion, and scope differences may explain the gap."})
    vdt_supply = _vdt_supply_processes(bundle)
    evidence_processes = set(_text(sources, "process")).union(_text(source_costs, "process"))
    outside_scope = sorted(process for process in evidence_processes if process and process not in vdt_supply)
    checks.append({
        "Check": "Hydrogen source evidence is within VDT supply scope",
        "Status": "Pass" if vdt_supply and not outside_scope else "Warning",
        "Actual": len(outside_scope),
        "Expected": 0.0,
        "Difference": float(len(outside_scope)),
        "Detail": (
            f"{len(evidence_processes)} active Hydrogen source process(es) have selected-run VDT supply evidence."
            if vdt_supply and not outside_scope else
            (
                "No VDT-derived supply scope is available; Hydrogen source evidence cannot be verified."
                if not vdt_supply else
                "Hydrogen source evidence outside the VDT-derived supply boundary: "
                + ", ".join(outside_scope)
            )
        ),
    })
    source_field = hydrogen_sankey_aggregation_field(source_aggregation)
    if source_field is not None:
        evidence = pd.concat([sources, source_costs], ignore_index=True)
        missing_labels = sorted(set(_text(evidence, "process").loc[_text(evidence, source_field).eq("")]))
        checks.append({
            "Check": f"Hydrogen source labels support {source_aggregation}",
            "Status": "Pass" if not missing_labels else "Warning",
            "Actual": len(missing_labels),
            "Expected": 0.0,
            "Difference": float(len(missing_labels)),
            "Detail": (
                "All active Hydrogen source evidence has the selected presentation label."
                if not missing_labels else
                "Missing presentation label; VDT-scoped evidence remains included for: "
                + ", ".join(missing_labels)
            ),
        })
    view = result.view_summary
    if not view.empty:
        selected = view[_text(view, "view").eq(UPSTREAM_SUPPLY_VIEW)]
        if not selected.empty:
            row = selected.iloc[0]
            checks.append({"Check": "Full Supply-chain view balance", "Status": "Pass" if abs(float(row.get("reconciliation_difference_m", 0.0))) <= _TOLERANCE else "Warning", "Actual": float(row.get("source_cost_m", 0.0)), "Expected": float(row.get("assigned_cost_m", 0.0)) + float(row.get("unassigned_source_cost_m", 0.0)), "Difference": float(row.get("reconciliation_difference_m", 0.0)), "Detail": "All-source view balance is reported independently of the Hydrogen subset."})
    return pd.DataFrame(checks)


def build_hydrogen_proof(
    result: ResourceAllocationResult,
    bundle,
    period: int,
    *,
    source_aggregation: str = DEFAULT_HYDROGEN_SOURCE_AGGREGATION,
    use_aggregation: str = DEFAULT_HYDROGEN_USE_AGGREGATION,
) -> HydrogenProof:
    """Build exact Hydrogen evidence with a configurable Sankey presentation."""
    hydrogen_sankey_aggregation_field(source_aggregation)
    hydrogen_sankey_aggregation_field(use_aggregation)
    source_costs = _source_costs(result, bundle, period)
    sources = _source_outputs(result, bundle, period, source_costs)
    ledger = _hydrogen_ledger(result, period)
    users, excluded_flows, sector_shares, basis_shares = _users(
        result, bundle, period, source_costs, ledger
    )
    physical_links, financial_links = _links(
        sources,
        source_costs,
        users,
        ledger,
        source_aggregation=source_aggregation,
        use_aggregation=use_aggregation,
    )
    vintage = result.vintage_allocations.copy()
    if not vintage.empty:
        vintage = vintage[
            _text(vintage, "view").eq(UPSTREAM_SUPPLY_VIEW)
            & _text(vintage, "resource_pool").eq(HYDROGEN_POOL)
            & pd.to_numeric(vintage.get("period"), errors="coerce").eq(period)
        ].copy()
        selected_source_rows = set(source_costs.get("source_row_id", pd.Series(dtype=str)).astype(str))
        vintage = vintage[vintage.get("source_row_id", pd.Series("", index=vintage.index)).astype(str).isin(selected_source_rows)]
    diagnostics = result.diagnostics.copy()
    if not diagnostics.empty and "period" in diagnostics.columns:
        diagnostics = diagnostics[pd.to_numeric(diagnostics["period"], errors="coerce").eq(period)].copy()
    checks = _checks(
        source_costs,
        sector_shares,
        sources,
        users,
        financial_links,
        result,
        bundle,
        period,
        source_aggregation,
    )
    return HydrogenProof(str(bundle.run_id), int(period), sources, source_costs, users, excluded_flows, sector_shares, financial_links, physical_links, vintage, checks, diagnostics, basis_shares)


def hydrogen_zoom_rows(proof: HydrogenProof, process: str, process_dukes: str) -> pd.DataFrame:
    """Replay one source process's cost rows into one DUKES classification.

    Each row uses the DUKES share of its own share-basis year, so investment
    follows the use mix of the year the capacity was built.
    """
    columns = [
        "source_row_id", "attribute", "vintage", "vintage_basis", "share_basis_year",
        "share_basis", "assigned_cost_m", "dukes_share", "allocated_to_selected_dukes_m",
    ]
    rows = proof.source_costs[_text(proof.source_costs, "process").eq(process)].copy()
    if rows.empty:
        return _empty(columns)
    rows["share_basis_year"] = _basis_years(rows).astype("Int64")
    shares = proof.basis_shares[proof.basis_shares["process_dukes"].eq(process_dukes)].copy()
    shares["share_basis_year"] = shares["share_basis_year"].astype("Int64")
    rows = rows.merge(
        shares[["share_basis_year", "dukes_share"]], on="share_basis_year", how="left"
    )
    rows["dukes_share"] = _numeric(rows, "dukes_share")
    rows["allocated_to_selected_dukes_m"] = _numeric(rows, "assigned_cost_m") * rows["dukes_share"]
    return rows.loc[:, columns].reset_index(drop=True)


def hydrogen_proof_filename(run_id: str, period: int) -> str:
    invalid = '<>:"/\\|?*'
    safe = "".join("_" if character in invalid else character for character in str(run_id)).strip()
    return f"{safe or 'selected_run'}_{period}_hydrogen_attribution_proof.xlsx"


def hydrogen_proof_workbook_bytes(proof: HydrogenProof, *, price_year: int | None = None) -> bytes:
    """Serialize transparent evidence into a formatted multi-sheet workbook."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("XLSX export requires openpyxl.") from exc
    workbook = Workbook()
    readme = workbook.active
    readme.title = "Read me"
    readme.sheet_view.showGridLines = False
    rows = [
        ("Hydrogen Attribution Proof", f"{proof.run_id} — {proof.period}"),
        ("Price year", str(price_year) if price_year else "As reported by selected run"),
        ("Scope", "Net Hydrogen source injections (gross output minus same-process Hydrogen input) and full VDT supply-chain cost rows. Distribution and network processes with zero net injection are excluded from physical sources but their costs remain in the financial attribution."),
        ("Method", "Supply-chain costs of every process supplying Hydrogen, including the energy converters such as electrolysers use, enter the Hydrogen pool. The pool is split by observed Hydrogen demand. A process that turns Hydrogen into another pool's energy carries its share into that pool. Cost with no route to final use is shown as unallocated."),
        ("Formula", "Pool cost = supply. Each demand's share = demand TWh / total included Hydrogen TWh in the cost row's share-basis year. Supply = final use + conversion demand + unallocated."),
        ("Important", HYDROGEN_ATTRIBUTION_WARNING),
        ("Vintage policy", "Only source-reported Cost_Inv and Cost_Fom vintages are shown. No synthetic vintage split is created."),
    ]
    for row in rows:
        readme.append(row)
    readme.column_dimensions["A"].width = 24
    readme.column_dimensions["B"].width = 115
    for row in readme.iter_rows():
        row[0].font = Font(bold=True, color="FFFFFF")
        row[0].fill = PatternFill("solid", fgColor="1F4E78")
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    readme.freeze_panes = "A2"

    sheets = [
        ("Proof checks", proof.checks),
        ("Net hydrogen sources", proof.sources),
        ("Source cost rows", proof.source_costs),
        ("All hydrogen users", proof.users),
        ("Excluded flows", proof.excluded_flows),
        ("DUKES shares", proof.sector_shares),
        ("Physical graph links", proof.physical_links),
        ("Financial graph links", proof.financial_links),
        ("Reported vintages", proof.vintage_detail),
        ("Diagnostics", proof.diagnostics),
    ]
    for title, frame in sheets:
        sheet = workbook.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        table = with_conversion_column_names(frame)
        sheet.append(list(table.columns))
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(wrap_text=True)
        for values in table.itertuples(index=False, name=None):
            sheet.append([None if pd.isna(value) else value for value in values])
        sheet.freeze_panes = "A2"
        for index, column in enumerate(table.columns, 1):
            width = min(max(len(str(column)) + 2, 14), 42)
            if not table.empty:
                maximum_length = table[column].fillna("").astype(str).str.len().max()
                if pd.notna(maximum_length):
                    width = min(max(width, min(42, int(maximum_length) + 2)), 42)
            sheet.column_dimensions[get_column_letter(index)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
