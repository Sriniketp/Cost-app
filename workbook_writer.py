"""Write the Resource & Fuel workbook as a traceable chain of live formulas.

Every derived cell is an Excel formula that points at the sheet before it, so a
reader can select a pooled figure and walk back to the raw model rows it came
from.  Values are pasted only on the raw evidence sheets, which are the start
of the chain, and on the pathway sheets, which are evidence rather than
calculation.

Blank cells are meaningful: they mean there is no evidence for that year, not
that the value is zero.  Interpolation never runs outside the years a series
was actually observed in.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO

import numpy as np
import pandas as pd

from cost_allocation import workbook_layout as layout
from cost_allocation.resource_allocation import (
    RESOURCE_FLOW_VIEW,
    with_conversion_column_names,
)
from cost_allocation.resource_export import (
    EXPORT_SECTORS,
    FUELEX_ROWS,
    RESOURCE_COST_ROWS,
    SOURCE_COST_FUELS,
    TWH_FUELS,
    PathwayExport,
    ResourceFuelExport,
)


TITLE_ROW = 1
SUBTITLE_ROW = 2
HEADER_ROW = 4
FIRST_DATA_ROW = 5

EXCEL_MAX_ROWS = 1_048_576

ALL_SECTORS = "All sectors"
ALL_EXCEPT_POWER = "All sectors except Power"
COMBINED_BIO = "Biomass & biofuels + Biomethane"
DISTRIBUTED_ELECTRICITY = "Electricity (distributed)"

# Label on sheets 40/41, the fuel-use row it divides by, and the source-cost row.
ALLOCATION_BASIS: tuple[tuple[str, str, str], ...] = (
    ("Synthetic fuel", "Synthetic fuel", "Ammonia & synthetics"),
    ("Biomass & biofuels", COMBINED_BIO, "Biomass & biofuels"),
    ("Hydrogen", "Hydrogen", "Hydrogen"),
    ("Natural gas", "Natural gas", "Natural gas"),
    ("Oil", "Oil", "Oil"),
)

_RAW_FUEL_LEAD = ("period", "sector", "resource_pool", "reference_fuel", "quantity_twh", "included")
_RAW_SOURCE_LEAD = ("period", "view", "resource_pool", "source_fuel", "source_cost_m")
_SOURCE_ASSIGNMENT_LEAD = (
    "source_row_id", "period", "view", "process", "process_name", "attribute",
    "cost_commodity", "cost_value_m", "resource_pool", "assignment_share",
    "assigned_cost_m", "assignment_status", "assignment_reason",
)
_RAW_POWER_LEAD = ("period", "process", "attribute", "power_core_cost_m")
_FLOW_PROVENANCE_LEAD = (
    "period", "commodity", "commodity_name", "process", "process_name", "direction",
    "reported_value", "unit", "vd_row_count", "timeslice_mode", "vdt_match",
    "provenance_note",
)
_COVERAGE_LEAD = (
    "period", "pool_kind", "commodity_name", "commodity", "resource_pool", "counted_twh",
    "excluded_twh", "coverage_status", "coverage_reason",
)

_MONEY_FORMAT = "#,##0.000"
_ENERGY_FORMAT = "#,##0.0000"
_SHARE_FORMAT = "0.000%"


@dataclass(frozen=True)
class WorkbookEvidence:
    """Raw application tables that are reproduced in the workbook unchanged."""

    pool_summary: pd.DataFrame = field(default_factory=pd.DataFrame)
    commodity_mapping: pd.DataFrame = field(default_factory=pd.DataFrame)
    model_connections: pd.DataFrame = field(default_factory=pd.DataFrame)
    commodity_costs: pd.DataFrame = field(default_factory=pd.DataFrame)
    vintage_allocations: pd.DataFrame = field(default_factory=pd.DataFrame)
    allocations: pd.DataFrame = field(default_factory=pd.DataFrame)
    diagnostics: pd.DataFrame = field(default_factory=pd.DataFrame)
    flow_provenance: pd.DataFrame = field(default_factory=pd.DataFrame)
    unit_conversion: pd.DataFrame = field(default_factory=pd.DataFrame)
    pooling_rules: pd.DataFrame = field(default_factory=pd.DataFrame)
    commodity_coverage: pd.DataFrame = field(default_factory=pd.DataFrame)
    source_assignments: pd.DataFrame = field(default_factory=pd.DataFrame)


@dataclass(frozen=True)
class TableAnchor:
    """Where one written table's data lives, so formulas can point at it."""

    sheet: str
    first_row: int
    last_row: int
    columns: dict[str, int]

    def column_range(self, name: str) -> str:
        return layout.column_range_ref(
            self.sheet, self.columns[name], self.first_row, self.last_row
        )


# ---------------------------------------------------------------------------
# Value helpers
# ---------------------------------------------------------------------------

def _clean(value: object) -> object:
    """Return a value openpyxl can write, mapping missing data to a blank cell."""
    if value is None:
        return None
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    if isinstance(value, (int, float, str)):
        return value
    return str(value)


def _ordered_columns(frame: pd.DataFrame, lead: tuple[str, ...]) -> list[str]:
    present = [name for name in lead if name in frame.columns]
    return present + [name for name in frame.columns if name not in present]


def _milestone_map(
    frame: pd.DataFrame, key_columns: list[str], value_column: str
) -> dict[tuple, list[int]]:
    """Return the model years each series was actually observed in."""
    required = {*key_columns, value_column, "period"}
    if frame is None or frame.empty or not required.issubset(frame.columns):
        return {}
    rows = frame.dropna(subset=["period", value_column])
    observed: dict[tuple, list[int]] = {}
    for key, group in rows.groupby(key_columns, dropna=False, sort=False):
        identity = key if isinstance(key, tuple) else (key,)
        observed[identity] = sorted({int(period) for period in group["period"]})
    return observed


def _has_fuel_use(export: ResourceFuelExport, fuel: str, sector: str) -> bool:
    """Return whether one reporting row has any milestone evidence."""
    rows = export.milestone_fuel_use
    if rows.empty or not {"sector", "reference_fuel"}.issubset(rows.columns):
        return False
    selected = rows["sector"].eq(sector)
    if fuel == COMBINED_BIO:
        selected &= rows["reference_fuel"].isin(("Biomass & biofuels", "Biomethane"))
    else:
        selected &= rows["reference_fuel"].eq(fuel)
    return bool(selected.any())


def _has_source_cost(export: ResourceFuelExport, fuel: str) -> bool:
    """Return whether one source-fuel row has any milestone evidence."""
    rows = export.milestone_source_costs
    return bool(
        not rows.empty
        and "source_fuel" in rows.columns
        and rows["source_fuel"].eq(fuel).any()
    )


def _interpolation_formula(milestones: list[int], year: int, reference) -> str | None:
    """Return a straight-line formula between the two bracketing model years.

    Returns ``None`` outside the observed range, which leaves the cell blank
    and reproduces the no-extrapolation rule the application applies.
    """
    if not milestones or year < milestones[0] or year > milestones[-1]:
        return None
    if year in milestones:
        return f"={reference(year)}"
    lower = max(period for period in milestones if period < year)
    upper = min(period for period in milestones if period > year)
    low_ref, high_ref = reference(lower), reference(upper)
    return f"={low_ref}+({high_ref}-{low_ref})*({year}-{lower})/({upper}-{lower})"


# ---------------------------------------------------------------------------
# Sheet scaffolding
# ---------------------------------------------------------------------------

def _styles():
    from openpyxl.styles import Alignment, Font, PatternFill

    return {
        "title": (Font(bold=True, size=14, color="FFFFFF"), PatternFill("solid", fgColor="17365D")),
        "header": (Font(bold=True), PatternFill("solid", fgColor="D9EAF7")),
        "stage": (Font(bold=True, italic=True), PatternFill("solid", fgColor="EEF4FA")),
        "wrap": Alignment(wrap_text=True, vertical="top"),
        "centre": Alignment(horizontal="center"),
    }


def _add_sheet(workbook, name: str, title: str, subtitle: str, width: int):
    styles = _styles()
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    title_font, title_fill = styles["title"]
    sheet.cell(TITLE_ROW, 1, title).font = title_font
    sheet.cell(TITLE_ROW, 1).fill = title_fill
    sheet.merge_cells(
        start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=max(width, 2)
    )
    sheet.cell(SUBTITLE_ROW, 1, subtitle).alignment = styles["wrap"]
    sheet.merge_cells(
        start_row=SUBTITLE_ROW, start_column=1, end_row=SUBTITLE_ROW, end_column=max(width, 2)
    )
    sheet.row_dimensions[SUBTITLE_ROW].height = 30
    return sheet


def _write_headers(sheet, headers: list[object]) -> None:
    styles = _styles()
    header_font, header_fill = styles["header"]
    for index, heading in enumerate(headers, 1):
        cell = sheet.cell(HEADER_ROW, index, heading)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = styles["centre"]


def _size_columns(sheet, label_widths: list[int], year_count: int, start: int) -> None:
    for index, width in enumerate(label_widths, 1):
        sheet.column_dimensions[layout.column_letter(index)].width = width
    for offset in range(year_count):
        sheet.column_dimensions[layout.column_letter(start + offset)].width = 13


def _is_missing(value: object) -> bool:
    if value is None:
        return True
    try:
        return bool(pd.isna(value))
    except (TypeError, ValueError):
        return False


def _column_values(series: pd.Series) -> list:
    """Return one column as native Python values, with missing data as blanks.

    Converting a column at a time keeps the pathway sheets, which can run to
    hundreds of thousands of rows, from paying a per-cell conversion cost.
    """
    if pd.api.types.is_bool_dtype(series):
        return series.fillna(False).astype(bool).tolist()
    values = series.tolist()
    return [
        None if _is_missing(value)
        else value if isinstance(value, (bool, int, float, str))
        else str(value)
        for value in values
    ]


def _write_value_table(
    workbook,
    name: str,
    title: str,
    subtitle: str,
    frame: pd.DataFrame,
    *,
    lead: tuple[str, ...] = (),
    warnings: list[str] | None = None,
) -> TableAnchor:
    """Write one raw table with plain-English headings and no formulas."""
    frame = pd.DataFrame() if frame is None else with_conversion_column_names(frame)
    columns = _ordered_columns(frame, lead)
    limit = EXCEL_MAX_ROWS - FIRST_DATA_ROW - 1
    if len(frame) > limit:
        if warnings is not None:
            warnings.append(
                f"{name} holds {len(frame):,} rows, so only the first {limit:,} are written. "
                "Excel cannot hold more rows than that on one sheet."
            )
        frame = frame.head(limit)
    sheet = _add_sheet(workbook, name, title, subtitle, max(len(columns), 2))
    _write_headers(sheet, layout.readable_headers(columns) or ["No data"])
    if len(frame) and columns:
        prepared = [_column_values(frame[column]) for column in columns]
        for values in zip(*prepared):
            sheet.append(list(values))
    else:
        sheet.cell(
            FIRST_DATA_ROW, 1,
            "No rows are available for this table in the selected run.",
        )
    sheet.freeze_panes = sheet.cell(FIRST_DATA_ROW, 1).coordinate
    for index, column in enumerate(columns, 1):
        letter = layout.column_letter(index)
        sheet.column_dimensions[letter].width = min(max(len(str(column)) + 8, 14), 32)
    last_row = max(FIRST_DATA_ROW + len(frame) - 1, FIRST_DATA_ROW)
    return TableAnchor(
        sheet=name,
        first_row=FIRST_DATA_ROW,
        last_row=last_row,
        columns={column: index for index, column in enumerate(columns, 1)},
    )


def _write_grid(
    workbook,
    name: str,
    title: str,
    subtitle: str,
    labels: list[str],
    years: tuple[int, ...],
) -> tuple:
    """Create a year-column grid and return the sheet plus its first year column."""
    first_year_column = len(labels) + 1
    sheet = _add_sheet(workbook, name, title, subtitle, first_year_column + len(years) - 1)
    _write_headers(sheet, [*labels, *years])
    _size_columns(sheet, [30 if index == 0 else 24 for index in range(len(labels))], len(years), first_year_column)
    sheet.freeze_panes = sheet.cell(FIRST_DATA_ROW, first_year_column).coordinate
    return sheet, first_year_column


def _set_formula(sheet, row: int, column: int, formula: str | None, number_format: str) -> None:
    cell = sheet.cell(row, column, formula)
    cell.number_format = number_format


def _mixed(column: int, row: int, *, lock_column: bool = False, lock_row: bool = False) -> str:
    """Return an A1 reference with independently locked column and row."""
    letter = layout.column_letter(column)
    return f"{'$' if lock_column else ''}{letter}{'$' if lock_row else ''}{row}"


def _criteria(anchor: TableAnchor, pairs: list[tuple[str, str]]) -> str:
    return ", ".join(
        f"{anchor.column_range(column)}, {criterion}" for column, criterion in pairs
    )


def _guarded_sumifs(anchor: TableAnchor, value_column: str, pairs: list[tuple[str, str]]) -> str:
    """Return a SUMIFS that stays blank when no evidence row matches at all."""
    criteria = _criteria(anchor, pairs)
    return (
        f"=IF(COUNTIFS({criteria})=0,\"\","
        f"SUMIFS({anchor.column_range(value_column)}, {criteria}))"
    )


# ---------------------------------------------------------------------------
# Guide sheets
# ---------------------------------------------------------------------------

def _write_read_me(workbook, export: ResourceFuelExport) -> None:
    styles = _styles()
    sheet = _add_sheet(
        workbook,
        layout.READ_ME_SHEET,
        f"{export.run_id} — Resource and fuel costs, full working",
        "This workbook shows the whole calculation, from the model's own rows through to the "
        "pooled totals. Every calculated cell is a live formula pointing at the sheet before it.",
        4,
    )
    price_year = export.price_year
    basis = f"{price_year} prices" if price_year else "the model's own price basis"
    notes = [
        ("Selected run", str(export.run_id)),
        ("Money basis", f"All £ figures are in {basis}, per year."),
        ("Years covered", f"{export.years[0]}–{export.years[-1]}" if export.years else "None"),
        ("Model years with evidence", ", ".join(str(year) for year in export.milestone_years) or "None"),
        ("How to trace a number", "Select any calculated cell and press Ctrl+[ to jump to the cells it is built from. Press Ctrl+] to see what uses it."),
        ("Formulas recalculate on open", "Cells are stored as formulas, not as pasted numbers. Excel calculates them when the file opens, so a value viewer that does not calculate will show them as empty."),
        ("Blank is not zero", "A blank calculated cell means there is no evidence for that year. Straight-line interpolation is never extended beyond the model years a series was actually observed in."),
        ("Rows hidden by default", "Complete reporting grids include some fuel and sector combinations with no evidence anywhere in this run. Those rows are retained with their formulas but hidden to keep the sheets readable; use Excel's Unhide Rows command to inspect them."),
        ("Do not add the views together", "Flow-cost Attribution and Full Supply-chain Attribution describe the same costs in different ways. They overlap and must never be summed, and neither may be added to the direct-cost totals on the Overview."),
        ("Commodity costs", "Cost_Com is disclosed on its own sheet and is deliberately excluded from every formula here."),
        ("Power and electricity", "The power sector receives no distributed electricity cost. That would allocate the power system's costs back to itself."),
    ]
    row = HEADER_ROW
    _write_headers(sheet, ["Item", "What it means"])
    row = FIRST_DATA_ROW
    for label, detail in notes:
        sheet.cell(row, 1, label).font = styles["header"][0]
        sheet.cell(row, 2, detail).alignment = styles["wrap"]
        row += 1

    row += 1
    stage_font, stage_fill = styles["stage"]
    heading = sheet.cell(row, 1, "Sheets, in calculation order")
    heading.font = stage_font
    heading.fill = stage_fill
    row += 1
    for column, title in enumerate(("Sheet", "Stage", "What it holds"), 1):
        cell = sheet.cell(row, column, title)
        cell.font = styles["header"][0]
        cell.fill = styles["header"][1]
    row += 1
    for spec in layout.SHEET_REGISTRY:
        sheet.cell(row, 1, spec.name)
        sheet.cell(row, 2, spec.stage)
        sheet.cell(row, 3, spec.description).alignment = styles["wrap"]
        row += 1
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 96


def _write_glossary(workbook) -> None:
    styles = _styles()
    sheet = _add_sheet(
        workbook,
        layout.GLOSSARY_SHEET,
        "Glossary",
        "Every abbreviation, unit and model term that survives anywhere in this workbook.",
        2,
    )
    _write_headers(sheet, ["Term", "Meaning"])
    for offset, (term, meaning) in enumerate(layout.GLOSSARY_ENTRIES):
        sheet.cell(FIRST_DATA_ROW + offset, 1, term)
        sheet.cell(FIRST_DATA_ROW + offset, 2, meaning).alignment = styles["wrap"]
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 110
    sheet.freeze_panes = sheet.cell(FIRST_DATA_ROW, 1).coordinate


_POOLING_NOTES = (
    ("One row per commodity", "Each commodity code is looked up by its exact code in the curated pooling sheet below. There are no ordered rules and nothing is inferred from labels."),
    ("Only the fuel decides", "The pool belongs to the commodity itself. The process consuming or producing it never changes which pool it belongs to."),
    ("One pool per fuel", "A commodity belongs to at most one resource pool, so no fuel use or source cost can be counted twice."),
    ("Not pooled is a reviewed answer", "A commodity listed without a pool was reviewed and deliberately left out of every pool. Its use is not counted, and it does not leave a producer's costs unallocated."),
    ("Demand pools stay outside the split", "Demand services grouped into a demand pool are listed for reference only. They are not fuels and never enter a fuel-cost split."),
    ("Anything absent is named", "An energy commodity missing from the sheet is given no pool. It is listed on the coverage sheet as needing attention rather than quietly dropped."),
    ("Where the result is shown", "The pool each commodity ended up in is on the fuel type mapping sheet. Which fuels count towards the split is on the coverage sheet."),
)


def _write_pooling_rules(workbook, rules: pd.DataFrame) -> None:
    """Show the curated pool sheet and how it is applied."""
    styles = _styles()
    sheet = _add_sheet(
        workbook,
        layout.POOLING_RULES_SHEET,
        "How each fuel is grouped into a resource pool",
        "Costs are attributed at the level of a resource pool. This curated sheet assigns each "
        "commodity code to one pool, and it is the whole of that decision.",
        4,
    )
    _write_headers(sheet, ["Item", "What it means"])
    row = FIRST_DATA_ROW
    for label, detail in _POOLING_NOTES:
        sheet.cell(row, 1, label).font = styles["header"][0]
        sheet.cell(row, 2, detail).alignment = styles["wrap"]
        row += 1

    row += 1
    stage_font, stage_fill = styles["stage"]
    heading = sheet.cell(row, 1, "The curated pooling sheet, one row per commodity")
    heading.font = stage_font
    heading.fill = stage_fill
    row += 1
    columns = list(rules.columns) if rules is not None and not rules.empty else []
    headings = layout.readable_headers(columns) or ["No data"]
    for index, heading_text in enumerate(headings, 1):
        cell = sheet.cell(row, index, heading_text)
        cell.font = styles["header"][0]
        cell.fill = styles["header"][1]
    row += 1
    if columns:
        for values in zip(*(_column_values(rules[column]) for column in columns)):
            for index, value in enumerate(values, 1):
                sheet.cell(row, index, value)
            row += 1
    else:
        sheet.cell(row, 1, "No pooling sheet is available for this run.")
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 96
    sheet.column_dimensions["C"].width = 30
    sheet.column_dimensions["D"].width = 30


def _write_unit_conversion(workbook, conversion: pd.DataFrame, provenance: TableAnchor) -> None:
    """Convert each reported unit to TWh in the workbook, as a live formula."""
    sheet = _add_sheet(
        workbook,
        layout.UNIT_CONVERSION_SHEET,
        "Turning reported units into TWh",
        "Nothing is added up until it is in TWh. One TWh is 3.6 PJ, so a PJ figure is multiplied "
        "by 0.2778. Each row is one model year, commodity and reported unit. Its total comes "
        "from matching input flows that belong to the run's model connections.",
        8,
    )
    headings = [
        layout.readable_header(column)
        for column in ("period", "commodity", "commodity_name", "unit", "unit_factor")
    ] + ["Reported input total", layout.readable_header("quantity_twh"), "Conversion status"]
    _write_headers(sheet, headings)
    conversion = pd.DataFrame() if conversion is None else conversion
    if conversion.empty:
        sheet.cell(FIRST_DATA_ROW, 1, "No reported energy units are available for this run.")
        return
    for offset, (_, entry) in enumerate(conversion.iterrows()):
        row = FIRST_DATA_ROW + offset
        period = int(entry["period"])
        code = str(entry["commodity"])
        unit = str(entry["unit"])
        sheet.cell(row, 1, period)
        sheet.cell(row, 2, code)
        sheet.cell(row, 3, str(entry["commodity_name"]))
        sheet.cell(row, 4, unit)
        factor = entry["unit_factor"]
        if pd.notna(factor):
            sheet.cell(row, 5, float(factor)).number_format = "0.0000"
        _set_formula(
            sheet, row, 6,
            "=SUMIFS({}, {}, {}, {}, {}, {}, {}, {}, \"IN\", {}, \"Yes\")".format(
                provenance.column_range("reported_value"),
                provenance.column_range("period"), period,
                provenance.column_range("commodity"), _quoted(code),
                provenance.column_range("unit"), _quoted(unit),
                provenance.column_range("direction"),
                provenance.column_range("vdt_match"),
            ),
            _ENERGY_FORMAT,
        )
        reported = _mixed(6, row)
        _set_formula(
            sheet, row, 7,
            f"=IF({_mixed(5, row)}=\"\",\"\",{reported}*{_mixed(5, row)})",
            _ENERGY_FORMAT,
        )
        sheet.cell(row, 8, str(entry["conversion_status"]))
    sheet.freeze_panes = sheet.cell(FIRST_DATA_ROW, 1).coordinate
    for index, width in enumerate((14, 22, 30, 18, 22, 20, 22, 28), 1):
        sheet.column_dimensions[layout.column_letter(index)].width = width


def _write_source_cost_summary(
    workbook,
    summary: pd.DataFrame,
    assignments: TableAnchor,
    *,
    has_assignment_rows: bool,
) -> TableAnchor:
    """Total source costs from process-level assignment rows with live formulas."""
    columns = list(_RAW_SOURCE_LEAD)
    summary = pd.DataFrame(columns=columns) if summary is None else summary.copy()
    for column in columns:
        if column not in summary.columns:
            summary[column] = pd.NA
    summary = summary.loc[:, columns]
    sheet = _add_sheet(
        workbook,
        layout.RAW_SOURCE_COST_SHEET,
        "Source costs totalled by model year, view and resource pool",
        "One row is one model year, attribution view and resource pool. Source cost is a live "
        "SUMIFS over the process-level assignments on sheet 09. The source fuel type is the "
        "fixed reference-layout label used by the later milestone sheet.",
        len(columns),
    )
    _write_headers(sheet, layout.readable_headers(columns))
    for offset, entry in enumerate(summary.itertuples(index=False)):
        row = FIRST_DATA_ROW + offset
        values = dict(zip(columns, entry))
        for column_index, column in enumerate(columns[:-1], 1):
            sheet.cell(row, column_index, _clean(values[column]))
        if has_assignment_rows:
            pairs = [
                ("period", _mixed(1, row)),
                ("view", _mixed(2, row)),
                ("resource_pool", _mixed(3, row)),
            ]
            _set_formula(
                sheet, row, 5,
                _guarded_sumifs(assignments, "assigned_cost_m", pairs),
                _MONEY_FORMAT,
            )
        else:
            sheet.cell(row, 5, _clean(values["source_cost_m"])).number_format = _MONEY_FORMAT
    if summary.empty:
        sheet.cell(FIRST_DATA_ROW, 1, "No source-cost pool rows are available for this run.")
    sheet.freeze_panes = sheet.cell(FIRST_DATA_ROW, 1).coordinate
    for index, width in enumerate((14, 30, 28, 24, 22), 1):
        sheet.column_dimensions[layout.column_letter(index)].width = width
    return TableAnchor(
        sheet=layout.RAW_SOURCE_COST_SHEET,
        first_row=FIRST_DATA_ROW,
        last_row=max(FIRST_DATA_ROW + len(summary) - 1, FIRST_DATA_ROW),
        columns={column: index for index, column in enumerate(columns, 1)},
    )


def _quoted(value: str) -> str:
    """Return a text literal safe to embed in an Excel criterion."""
    return '"{}"'.format(str(value).replace('"', '""'))


def _worked_example_choice(export: ResourceFuelExport):
    """Return the largest reported fuel, sector and year to follow as an example.

    Choosing the largest observed combination keeps the example meaningful.
    A run with no evidence at all returns nothing and the sheet says so.
    """
    rows = export.milestone_fuel_use
    if rows is None or rows.empty:
        return None
    required = {"period", "sector", "reference_fuel", "quantity_twh"}
    if not required.issubset(rows.columns) or not export.years:
        return None
    labels = {label: (use_fuel, source_fuel) for label, use_fuel, source_fuel in ALLOCATION_BASIS}
    candidates = rows[
        rows["reference_fuel"].isin(labels)
        & rows["sector"].isin(EXPORT_SECTORS)
        & pd.to_numeric(rows["period"], errors="coerce").isin(export.years)
    ].copy()
    candidates["quantity_twh"] = pd.to_numeric(candidates["quantity_twh"], errors="coerce")
    candidates = candidates[candidates["quantity_twh"].gt(0)]
    if candidates.empty:
        return None
    best = candidates.sort_values("quantity_twh", ascending=False).iloc[0]
    label = str(best["reference_fuel"])
    use_fuel, source_fuel = labels[label]
    return int(best["period"]), label, use_fuel, source_fuel, str(best["sector"])


def _add_worked_example(workbook):
    """Reserve the worked example sheet so it keeps its place in the order."""
    return _add_sheet(
        workbook,
        layout.WORKED_EXAMPLE_SHEET,
        "One allocated cost, followed all the way through",
        "Every value below is a live formula pointing at the sheet named beside it. Select a "
        "value and press Ctrl+[ to jump to the cells it is built from.",
        4,
    )


def _fill_worked_example(
    sheet,
    export: ResourceFuelExport,
    *,
    provenance: TableAnchor,
    use_column: int,
    use_rows: dict,
    cost_column: int,
    cost_rows: dict,
    share_column: int,
    share_rows: dict,
    allocated_column: int,
    allocated_rows: dict,
) -> None:
    styles = _styles()
    choice = _worked_example_choice(export)
    if choice is None:
        _write_headers(sheet, ["Step", "What happens"])
        sheet.cell(
            FIRST_DATA_ROW, 1,
            "This run has no reported fuel use, so there is no allocation to follow.",
        )
        sheet.column_dimensions["A"].width = 30
        sheet.column_dimensions["B"].width = 96
        return
    period, label, use_fuel, source_fuel, sector = choice
    offset = export.years.index(period)

    sheet.cell(SUBTITLE_ROW, 1).value = (
        f"Following {label} used by {sector} in {period}. Every value below is a live formula "
        "pointing at the sheet named beside it. Select a value and press Ctrl+[ to jump to the "
        "cells it is built from."
    )
    _write_headers(sheet, ["Step", "What happens", "Where to look", "Value"])

    sector_use = layout.cell_ref(
        layout.ANNUAL_FUEL_USE_SHEET, use_rows[(use_fuel, sector)], use_column + offset
    )
    total_use = layout.cell_ref(
        layout.ANNUAL_FUEL_USE_SHEET, use_rows[(use_fuel, ALL_SECTORS)], use_column + offset
    )
    pool_cost = layout.cell_ref(
        layout.ANNUAL_SOURCE_COST_SHEET, cost_rows[source_fuel], cost_column + offset
    )
    share = layout.cell_ref(
        layout.ALLOCATION_SHARE_SHEET, share_rows[(label, sector)], share_column + offset
    )
    allocated = layout.cell_ref(
        layout.ALLOCATED_COST_SHEET, allocated_rows[(label, sector)], allocated_column + offset
    )
    reported = (
        "=SUMIFS({}, {}, {}, {}, \"IN\", {}, \"Yes\")".format(
            provenance.column_range("vd_row_count"),
            provenance.column_range("period"), period,
            provenance.column_range("direction"),
            provenance.column_range("vdt_match"),
        )
    )
    steps = (
        ("1. Start with the model's own rows",
         f"Count the reported input rows behind {period}. Each one records how many result rows were combined.",
         layout.FLOW_PROVENANCE_SHEET, reported, "#,##0"),
        ("2. Put everything in TWh",
         "Reported units are converted before anything is added up. One TWh is 3.6 PJ.",
         layout.UNIT_CONVERSION_SHEET, None, None),
        ("3. Group the fuel into a pool",
         f"{label} is grouped by the curated pooling sheet, so costs are attributed at pool level.",
         layout.POOLING_RULES_SHEET, None, None),
        ("4. Add up this sector's use",
         f"How much {label} {sector} used in {period}.",
         layout.ANNUAL_FUEL_USE_SHEET, f"={sector_use}", _ENERGY_FORMAT),
        ("5. Add up everyone's use",
         f"How much {label} every sector used in {period}. This is what the cost is divided by.",
         layout.ANNUAL_FUEL_USE_SHEET, f"={total_use}", _ENERGY_FORMAT),
        ("6. Work out the share",
         "Share = this sector's use divided by everyone's use.",
         layout.ALLOCATION_SHARE_SHEET, f"={share}", _SHARE_FORMAT),
        ("7. Take the pool's source cost",
         f"The source cost recorded against {source_fuel} in {period}, before any attribution.",
         layout.ANNUAL_SOURCE_COST_SHEET, f"={pool_cost}", _MONEY_FORMAT),
        ("8. Multiply cost by share",
         "Allocated cost = source cost multiplied by the sector's share of use.",
         layout.ALLOCATED_COST_SHEET, f"={allocated}", _MONEY_FORMAT),
    )
    row = FIRST_DATA_ROW
    for step, description, where, formula, number_format in steps:
        sheet.cell(row, 1, step).font = styles["header"][0]
        sheet.cell(row, 2, description).alignment = styles["wrap"]
        sheet.cell(row, 3, where)
        if formula is not None:
            _set_formula(sheet, row, 4, formula, number_format)
        else:
            sheet.cell(row, 4, "See sheet")
        row += 1

    row += 1
    stage_font, stage_fill = styles["stage"]
    check = sheet.cell(row, 1, "Proof that step 8 is steps 6 and 7 multiplied")
    check.font = stage_font
    check.fill = stage_fill
    row += 1
    sheet.cell(row, 1, "Share × source cost").font = styles["header"][0]
    sheet.cell(row, 2, "Recalculated here from the two figures above.").alignment = styles["wrap"]
    sheet.cell(row, 3, layout.WORKED_EXAMPLE_SHEET)
    _set_formula(
        sheet, row, 4,
        f"=IF(OR({share}=\"\",{pool_cost}=\"\"),\"\",{share}*{pool_cost})",
        _MONEY_FORMAT,
    )
    recalculated = _mixed(4, row)
    row += 1
    sheet.cell(row, 1, "Difference from step 8").font = styles["header"][0]
    sheet.cell(row, 2, "Zero confirms the allocated cost is exactly cost times share.").alignment = styles["wrap"]
    sheet.cell(row, 3, layout.ALLOCATED_COST_SHEET)
    _set_formula(sheet, row, 4, f"=N({recalculated})-N({allocated})", _MONEY_FORMAT)

    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 76
    sheet.column_dimensions["C"].width = 28
    sheet.column_dimensions["D"].width = 20


# ---------------------------------------------------------------------------
# Milestone sheets
# ---------------------------------------------------------------------------

def _write_milestone_fuel_use(workbook, export: ResourceFuelExport, raw: TableAnchor):
    years = export.milestone_years
    sheet, first_year = _write_grid(
        workbook,
        layout.MILESTONE_FUEL_USE_SHEET,
        "Fuel use in each model year (TWh/yr)",
        "Each cell adds up the matching rows on the raw fuel use sheet. The fuel type and sector "
        "in columns A and B, and the model year in the heading, are the search terms used.",
        ["Reported fuel type", "DESNZ sector"],
        years,
    )
    rows: dict[tuple[str, str], int] = {}
    row = FIRST_DATA_ROW
    for fuel in TWH_FUELS:
        for sector in EXPORT_SECTORS:
            sheet.cell(row, 1, fuel)
            sheet.cell(row, 2, sector)
            for offset, year in enumerate(years):
                column = first_year + offset
                pairs = [
                    ("period", _mixed(column, HEADER_ROW, lock_row=True)),
                    ("sector", _mixed(2, row, lock_column=True)),
                    ("reference_fuel", _mixed(1, row, lock_column=True)),
                    ("included", "TRUE"),
                ]
                _set_formula(
                    sheet, row, column,
                    _guarded_sumifs(raw, "quantity_twh", pairs), _ENERGY_FORMAT,
                )
            rows[(fuel, sector)] = row
            if not _has_fuel_use(export, fuel, sector):
                sheet.row_dimensions[row].hidden = True
            row += 1
    return sheet, first_year, rows


def _write_milestone_source_costs(workbook, export: ResourceFuelExport, raw: TableAnchor):
    years = export.milestone_years
    sheet, first_year = _write_grid(
        workbook,
        layout.MILESTONE_SOURCE_COST_SHEET,
        "Source cost in each model year (£m/yr)",
        "Each cell adds up the matching rows on the raw source cost sheet, keeping only the "
        f"'{RESOURCE_FLOW_VIEW}' view that flow-cost attribution is built from.",
        ["Source fuel type"],
        years,
    )
    rows: dict[str, int] = {}
    for offset_row, fuel in enumerate(SOURCE_COST_FUELS):
        row = FIRST_DATA_ROW + offset_row
        sheet.cell(row, 1, fuel)
        for offset, year in enumerate(years):
            column = first_year + offset
            pairs = [
                ("period", _mixed(column, HEADER_ROW, lock_row=True)),
                ("source_fuel", _mixed(1, row, lock_column=True)),
                ("view", layout.text_literal(RESOURCE_FLOW_VIEW)),
            ]
            _set_formula(
                sheet, row, column,
                _guarded_sumifs(raw, "source_cost_m", pairs), _MONEY_FORMAT,
            )
        rows[fuel] = row
        if not _has_source_cost(export, fuel):
            sheet.row_dimensions[row].hidden = True
    return sheet, first_year, rows


# ---------------------------------------------------------------------------
# Annual sheets
# ---------------------------------------------------------------------------

def _write_annual_fuel_use(workbook, export: ResourceFuelExport, milestone_columns, milestone_rows):
    years = export.years
    observed = _milestone_map(
        export.milestone_fuel_use, ["sector", "reference_fuel"], "quantity_twh"
    )
    sheet, first_year = _write_grid(
        workbook,
        layout.ANNUAL_FUEL_USE_SHEET,
        "Fuel use in every year (TWh/yr)",
        "Model years are taken straight from the milestone sheet. Years in between are a straight "
        "line drawn between the two nearest model years. Years outside the evidence stay blank.",
        ["Reported fuel type", "DESNZ sector"],
        years,
    )
    rows: dict[tuple[str, str], int] = {}
    row = FIRST_DATA_ROW
    for fuel in TWH_FUELS:
        block_start = row
        for sector in EXPORT_SECTORS:
            sheet.cell(row, 1, fuel)
            sheet.cell(row, 2, sector)
            milestones = observed.get((sector, fuel), [])
            source_row = milestone_rows[(fuel, sector)]

            def reference(period: int, source_row: int = source_row) -> str:
                return layout.cell_ref(
                    layout.MILESTONE_FUEL_USE_SHEET, source_row, milestone_columns[period]
                )

            for offset, year in enumerate(years):
                formula = _interpolation_formula(milestones, year, reference)
                _set_formula(sheet, row, first_year + offset, formula, _ENERGY_FORMAT)
            rows[(fuel, sector)] = row
            if not milestones:
                sheet.row_dimensions[row].hidden = True
            row += 1
        block_end = row - 1
        sheet.cell(row, 1, fuel)
        sheet.cell(row, 2, ALL_SECTORS).font = _styles()["header"][0]
        for offset in range(len(years)):
            column = first_year + offset
            span = f"{_mixed(column, block_start)}:{_mixed(column, block_end)}"
            _set_formula(sheet, row, column, f"=SUM({span})", _ENERGY_FORMAT)
        rows[(fuel, ALL_SECTORS)] = row
        row += 1
        if fuel == "Electricity":
            sheet.cell(row, 1, fuel)
            sheet.cell(row, 2, ALL_EXCEPT_POWER).font = _styles()["header"][0]
            total_row = rows[(fuel, ALL_SECTORS)]
            power_row = rows[(fuel, "Power")]
            for offset in range(len(years)):
                column = first_year + offset
                total = _mixed(column, total_row)
                power = _mixed(column, power_row)
                _set_formula(sheet, row, column, f"={total}-N({power})", _ENERGY_FORMAT)
            rows[(fuel, ALL_EXCEPT_POWER)] = row
            row += 1

    block_start = row
    for sector in EXPORT_SECTORS:
        sheet.cell(row, 1, COMBINED_BIO)
        sheet.cell(row, 2, sector)
        for offset in range(len(years)):
            column = first_year + offset
            letter = layout.column_letter(column)
            first = f"{letter}{rows[('Biomass & biofuels', sector)]}"
            second = f"{letter}{rows[('Biomethane', sector)]}"
            _set_formula(sheet, row, column, f"=N({first})+N({second})", _ENERGY_FORMAT)
        rows[(COMBINED_BIO, sector)] = row
        if not _has_fuel_use(export, COMBINED_BIO, sector):
            sheet.row_dimensions[row].hidden = True
        row += 1
    block_end = row - 1
    sheet.cell(row, 1, COMBINED_BIO)
    sheet.cell(row, 2, ALL_SECTORS).font = _styles()["header"][0]
    for offset in range(len(years)):
        column = first_year + offset
        span = f"{_mixed(column, block_start)}:{_mixed(column, block_end)}"
        _set_formula(sheet, row, column, f"=SUM({span})", _ENERGY_FORMAT)
    rows[(COMBINED_BIO, ALL_SECTORS)] = row
    return sheet, first_year, rows


def _write_annual_source_costs(workbook, export: ResourceFuelExport, milestone_columns, milestone_rows):
    years = export.years
    observed = _milestone_map(export.milestone_source_costs, ["source_fuel"], "source_cost_m")
    sheet, first_year = _write_grid(
        workbook,
        layout.ANNUAL_SOURCE_COST_SHEET,
        "Source cost in every year (£m/yr)",
        "Model years are taken straight from the milestone sheet. Years in between are a straight "
        "line drawn between the two nearest model years. Years outside the evidence stay blank.",
        ["Source fuel type"],
        years,
    )
    rows: dict[str, int] = {}
    for offset_row, fuel in enumerate(SOURCE_COST_FUELS):
        row = FIRST_DATA_ROW + offset_row
        sheet.cell(row, 1, fuel)
        milestones = observed.get((fuel,), [])
        source_row = milestone_rows[fuel]

        def reference(period: int, source_row: int = source_row) -> str:
            return layout.cell_ref(
                layout.MILESTONE_SOURCE_COST_SHEET, source_row, milestone_columns[period]
            )

        for offset, year in enumerate(years):
            formula = _interpolation_formula(milestones, year, reference)
            _set_formula(sheet, row, first_year + offset, formula, _MONEY_FORMAT)
        rows[fuel] = row
        if not milestones:
            sheet.row_dimensions[row].hidden = True
    total_row = FIRST_DATA_ROW + len(SOURCE_COST_FUELS)
    sheet.cell(total_row, 1, "Total").font = _styles()["header"][0]
    for offset in range(len(years)):
        column = first_year + offset
        span = f"{_mixed(column, FIRST_DATA_ROW)}:{_mixed(column, total_row - 1)}"
        _set_formula(sheet, total_row, column, f"=IF(COUNT({span})=0,\"\",SUM({span}))", _MONEY_FORMAT)
    rows["Total"] = total_row
    return sheet, first_year, rows


# ---------------------------------------------------------------------------
# Allocation sheets
# ---------------------------------------------------------------------------

def _write_allocation_shares(workbook, export: ResourceFuelExport, use_column, use_rows):
    years = export.years
    sheet, first_year = _write_grid(
        workbook,
        layout.ALLOCATION_SHARE_SHEET,
        "Each sector's share of total use of a fuel",
        "Share = the sector's own use of the fuel divided by every sector's use of it, in the same "
        "year. A share is left blank when nobody used that fuel, because there is nothing to divide by.",
        ["Fuel type", "DESNZ sector"],
        years,
    )
    rows: dict[tuple[str, str], int] = {}
    row = FIRST_DATA_ROW
    for label, use_fuel, _ in ALLOCATION_BASIS:
        for sector in EXPORT_SECTORS:
            sheet.cell(row, 1, label)
            sheet.cell(row, 2, sector)
            for offset in range(len(years)):
                column = first_year + offset
                use = layout.cell_ref(
                    layout.ANNUAL_FUEL_USE_SHEET, use_rows[(use_fuel, sector)],
                    use_column + offset,
                )
                total = layout.cell_ref(
                    layout.ANNUAL_FUEL_USE_SHEET, use_rows[(use_fuel, ALL_SECTORS)],
                    use_column + offset,
                )
                _set_formula(
                    sheet, row, column,
                    f"=IF(N({total})>0,N({use})/N({total}),\"\")", _SHARE_FORMAT,
                )
            rows[(label, sector)] = row
            if not _has_fuel_use(export, use_fuel, sector):
                sheet.row_dimensions[row].hidden = True
            row += 1
    return sheet, first_year, rows


def _write_allocated_costs(workbook, export, share_column, share_rows, cost_column, cost_rows):
    years = export.years
    sheet, first_year = _write_grid(
        workbook,
        layout.ALLOCATED_COST_SHEET,
        "Source cost shared out by use (£m/yr)",
        "Allocated cost = the source cost of the fuel multiplied by the sector's share of use. "
        "Bio costs use one combined pool of biomass, biofuels and biomethane.",
        ["Fuel type", "DESNZ sector"],
        years,
    )
    rows: dict[tuple[str, str], int] = {}
    blocks: dict[str, tuple[int, int]] = {}
    row = FIRST_DATA_ROW
    for label, _, source_fuel in ALLOCATION_BASIS:
        block_start = row
        for sector in EXPORT_SECTORS:
            sheet.cell(row, 1, label)
            sheet.cell(row, 2, sector)
            for offset in range(len(years)):
                share = layout.cell_ref(
                    layout.ALLOCATION_SHARE_SHEET, share_rows[(label, sector)],
                    share_column + offset,
                )
                cost = layout.cell_ref(
                    layout.ANNUAL_SOURCE_COST_SHEET, cost_rows[source_fuel],
                    cost_column + offset,
                )
                _set_formula(
                    sheet, row, first_year + offset,
                    f"=IF(OR({share}=\"\",{cost}=\"\"),\"\",{share}*{cost})", _MONEY_FORMAT,
                )
            rows[(label, sector)] = row
            if not _has_fuel_use(export, next(
                use_fuel for basis_label, use_fuel, _ in ALLOCATION_BASIS
                if basis_label == label
            ), sector):
                sheet.row_dimensions[row].hidden = True
            row += 1
        blocks[label] = (block_start, row - 1)
    return sheet, first_year, rows, blocks


def _write_electricity(workbook, export, power_raw, use_column, use_rows, allocated_column, allocated_rows):
    years = export.years
    sheet, first_year = _write_grid(
        workbook,
        layout.ELECTRICITY_COST_SHEET,
        "Electricity cost and how it reaches each sector (£m/yr unless stated)",
        "The power system's own core cost, less the fuel cost already charged to the power sector, "
        "gives a base cost. That base plus the power sector's fuel cost is then shared out by "
        "electricity use. The power sector receives none of it, which would otherwise be circular.",
        ["Measure", "DESNZ sector"],
        years,
    )
    labels = [
        "Power-system core cost",
        "Power sector fuel cost",
        "Electricity base cost at model years",
        "Electricity base cost, interpolated",
        "Distributed electricity cost",
        "Electricity Costs (£bn/yr)",
        "Electricity Costs + Power fuel costs (£bn/yr)",
    ]
    rows = {label: FIRST_DATA_ROW + index for index, label in enumerate(labels)}
    for label, row in rows.items():
        sheet.cell(row, 1, label)

    milestone_years = set(export.milestone_years)
    power_block = [allocated_rows[(label, "Power")] for label, _, _ in ALLOCATION_BASIS]
    for offset, year in enumerate(years):
        column = first_year + offset
        heading = _mixed(column, HEADER_ROW, lock_row=True)
        core_row = rows["Power-system core cost"]
        if year in milestone_years:
            _set_formula(
                sheet, core_row, column,
                _guarded_sumifs(power_raw, "power_core_cost_m", [("period", heading)]),
                _MONEY_FORMAT,
            )
        else:
            _set_formula(sheet, core_row, column, None, _MONEY_FORMAT)

        fuel_refs = ",".join(
            layout.cell_ref(layout.ALLOCATED_COST_SHEET, power_row, allocated_column + offset)
            for power_row in power_block
        )
        _set_formula(
            sheet, rows["Power sector fuel cost"], column,
            f"=IF(COUNT({fuel_refs})=0,\"\",SUM({fuel_refs}))", _MONEY_FORMAT,
        )

        core = _mixed(column, core_row)
        fuel = _mixed(column, rows["Power sector fuel cost"])
        _set_formula(
            sheet, rows["Electricity base cost at model years"], column,
            f"=IF(OR({core}=\"\",{fuel}=\"\"),\"\",{core}-{fuel})", _MONEY_FORMAT,
        )

    anchor_row = rows["Electricity base cost at model years"]
    base_years = [year for year in export.electricity_base_years if year in years]
    year_column = {year: first_year + index for index, year in enumerate(years)}

    def base_reference(period: int) -> str:
        return _mixed(year_column[period], anchor_row)

    for offset, year in enumerate(years):
        column = first_year + offset
        _set_formula(
            sheet, rows["Electricity base cost, interpolated"], column,
            _interpolation_formula(base_years, year, base_reference), _MONEY_FORMAT,
        )
        base = _mixed(column, rows["Electricity base cost, interpolated"])
        fuel = _mixed(column, rows["Power sector fuel cost"])
        _set_formula(
            sheet, rows["Distributed electricity cost"], column,
            f"=IF(OR({base}=\"\",{fuel}=\"\"),\"\",{base}+{fuel})", _MONEY_FORMAT,
        )
        distributed = _mixed(column, rows["Distributed electricity cost"])
        _set_formula(
            sheet, rows["Electricity Costs (£bn/yr)"], column,
            f"=IF({base}=\"\",\"\",{base}/1000)", _MONEY_FORMAT,
        )
        _set_formula(
            sheet, rows["Electricity Costs + Power fuel costs (£bn/yr)"], column,
            f"=IF({distributed}=\"\",\"\",{distributed}/1000)", _MONEY_FORMAT,
        )

    sector_rows: dict[str, int] = {}
    row = FIRST_DATA_ROW + len(labels) + 1
    block_start = row
    for sector in EXPORT_SECTORS:
        sheet.cell(row, 1, "Distributed to sector")
        sheet.cell(row, 2, sector)
        for offset in range(len(years)):
            column = first_year + offset
            if sector == "Power":
                cell = sheet.cell(row, column, 0.0)
                cell.number_format = _MONEY_FORMAT
                continue
            distributed = _mixed(column, rows["Distributed electricity cost"])
            use = layout.cell_ref(
                layout.ANNUAL_FUEL_USE_SHEET, use_rows[("Electricity", sector)],
                use_column + offset,
            )
            total = layout.cell_ref(
                layout.ANNUAL_FUEL_USE_SHEET, use_rows[("Electricity", ALL_EXCEPT_POWER)],
                use_column + offset,
            )
            _set_formula(
                sheet, row, column,
                f"=IF(AND({distributed}<>\"\",N({total})>0),{distributed}*N({use})/N({total}),\"\")",
                _MONEY_FORMAT,
            )
        sector_rows[sector] = row
        if sector != "Power" and not _has_fuel_use(export, "Electricity", sector):
            sheet.row_dimensions[row].hidden = True
        row += 1
    sheet.cell(
        row + 1, 1,
        "The power sector is held at zero on purpose. Charging the power system its own costs "
        "would count them twice.",
    )
    return sheet, first_year, rows, sector_rows, (block_start, row - 1)


def _write_resource_costs(workbook, export, allocated_column, allocated_rows, electricity_column, electricity_rows):
    years = export.years
    sheet, first_year = _write_grid(
        workbook,
        layout.RESOURCE_COST_SHEET,
        "Allocated cost for every sector and fuel (£m/yr)",
        "This sheet brings the allocated fuel costs and the distributed electricity cost together. "
        "Every cell points straight at the sheet it came from and recalculates nothing.",
        ["DESNZ sector", "Fuel type"],
        years,
    )
    rows: dict[tuple[str, str], int] = {}
    row = FIRST_DATA_ROW
    for sector in EXPORT_SECTORS:
        for fuel in RESOURCE_COST_ROWS:
            sheet.cell(row, 1, sector)
            sheet.cell(row, 2, fuel)
            for offset in range(len(years)):
                if fuel == DISTRIBUTED_ELECTRICITY:
                    reference = layout.cell_ref(
                        layout.ELECTRICITY_COST_SHEET, electricity_rows[sector],
                        electricity_column + offset,
                    )
                else:
                    reference = layout.cell_ref(
                        layout.ALLOCATED_COST_SHEET, allocated_rows[(fuel, sector)],
                        allocated_column + offset,
                    )
                _set_formula(
                    sheet, row, first_year + offset,
                    f"=IF({reference}=\"\",\"\",{reference})", _MONEY_FORMAT,
                )
            rows[(sector, fuel)] = row
            evidence_fuel = "Electricity" if fuel == DISTRIBUTED_ELECTRICITY else fuel
            if not _has_fuel_use(export, evidence_fuel, sector):
                sheet.row_dimensions[row].hidden = True
            row += 1
    return sheet, first_year, rows


_FUELEX_PARTS: dict[str, tuple[str, ...]] = {
    "Fuelex - alt": ("Synthetic fuel", "Biomass & biofuels", "Hydrogen"),
    "Fuelex - alt + elec": ("Synthetic fuel", "Biomass & biofuels", "Hydrogen", DISTRIBUTED_ELECTRICITY),
    "Fuelex - f/fuels": ("Natural gas", "Oil"),
    "Net Fuel Costs": ("Synthetic fuel", "Biomass & biofuels", "Hydrogen", "Natural gas", "Oil"),
    "Net Fuel Costs + elec": (
        "Synthetic fuel", "Biomass & biofuels", "Hydrogen", "Natural gas", "Oil",
        DISTRIBUTED_ELECTRICITY,
    ),
}


def _write_fuelex(workbook, export, resource_column, resource_rows):
    years = export.years
    sheet, first_year = _write_grid(
        workbook,
        layout.FUELEX_SHEET,
        "Grouped fuel-cost measures by sector (£bn/yr)",
        "Each measure adds together the relevant rows from the resource cost sheet and converts "
        "millions to billions. 'alt' means synthetic, bio and hydrogen fuels; 'f/fuels' means "
        "natural gas and oil.",
        ["DESNZ sector", "Measure"],
        years,
    )
    row = FIRST_DATA_ROW
    for sector in EXPORT_SECTORS:
        sector_has_evidence = any(
            _has_fuel_use(export, fuel, sector)
            for fuel in (
                "Synthetic fuel", COMBINED_BIO, "Hydrogen", "Natural gas", "Oil",
                "Electricity",
            )
        )
        for measure in FUELEX_ROWS:
            sheet.cell(row, 1, sector)
            sheet.cell(row, 2, measure)
            parts = _FUELEX_PARTS[measure]
            for offset in range(len(years)):
                references = ",".join(
                    layout.cell_ref(
                        layout.RESOURCE_COST_SHEET, resource_rows[(sector, fuel)],
                        resource_column + offset,
                    )
                    for fuel in parts
                )
                _set_formula(
                    sheet, row, first_year + offset,
                    f"=SUM({references})/1000", _MONEY_FORMAT,
                )
            if not sector_has_evidence:
                sheet.row_dimensions[row].hidden = True
            row += 1
    return sheet


# ---------------------------------------------------------------------------
# Audit sheets
# ---------------------------------------------------------------------------

def _write_checks(
    workbook, export, cost_column, cost_rows, allocated_column, allocated_blocks,
    electricity_column, electricity_rows, electricity_block, warnings: list[str],
):
    styles = _styles()
    years = export.years
    sheet, first_year = _write_grid(
        workbook,
        layout.CHECKS_SHEET,
        "Reconciliation and warnings",
        "Source cost should equal what was shared out to sectors plus what stayed unallocated. "
        "These rows prove that in the workbook itself rather than asserting it in prose.",
        ["Fuel type", "Measure"],
        years,
    )
    row = FIRST_DATA_ROW
    for label, _, source_fuel in ALLOCATION_BASIS:
        start, end = allocated_blocks[label]
        source_row = row
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, "Source cost (£m/yr)")
        for offset in range(len(years)):
            reference = layout.cell_ref(
                layout.ANNUAL_SOURCE_COST_SHEET, cost_rows[source_fuel], cost_column + offset
            )
            _set_formula(sheet, row, first_year + offset, f"=N({reference})", _MONEY_FORMAT)
        row += 1
        allocated_row = row
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, "Shared out to sectors (£m/yr)")
        for offset in range(len(years)):
            column = allocated_column + offset
            span = layout.range_ref(layout.ALLOCATED_COST_SHEET, start, column, end, column)
            _set_formula(sheet, row, first_year + offset, f"=SUM({span})", _MONEY_FORMAT)
        row += 1
        sheet.cell(row, 1, label)
        sheet.cell(row, 2, "Left unallocated (£m/yr)")
        for offset in range(len(years)):
            column = first_year + offset
            source = _mixed(column, source_row)
            allocated = _mixed(column, allocated_row)
            _set_formula(sheet, row, column, f"={source}-{allocated}", _MONEY_FORMAT)
        row += 1

    start, end = electricity_block
    source_row = row
    sheet.cell(row, 1, "Electricity")
    sheet.cell(row, 2, "Cost available to distribute (£m/yr)")
    for offset in range(len(years)):
        reference = layout.cell_ref(
            layout.ELECTRICITY_COST_SHEET, electricity_rows["Distributed electricity cost"],
            electricity_column + offset,
        )
        _set_formula(sheet, row, first_year + offset, f"=N({reference})", _MONEY_FORMAT)
    row += 1
    allocated_row = row
    sheet.cell(row, 1, "Electricity")
    sheet.cell(row, 2, "Shared out to sectors (£m/yr)")
    for offset in range(len(years)):
        column = electricity_column + offset
        span = layout.range_ref(layout.ELECTRICITY_COST_SHEET, start, column, end, column)
        _set_formula(sheet, row, first_year + offset, f"=SUM({span})", _MONEY_FORMAT)
    row += 1
    sheet.cell(row, 1, "Electricity")
    sheet.cell(row, 2, "Left unallocated (£m/yr)")
    for offset in range(len(years)):
        column = first_year + offset
        _set_formula(
            sheet, row, column,
            f"={_mixed(column, source_row)}-{_mixed(column, allocated_row)}", _MONEY_FORMAT,
        )
    row += 3

    heading = sheet.cell(row, 1, "Warnings raised while preparing this workbook")
    heading.font = styles["stage"][0]
    heading.fill = styles["stage"][1]
    row += 1
    for column, title in enumerate(("Check", "Detail", "Value"), 1):
        cell = sheet.cell(row, column, title)
        cell.font = styles["header"][0]
        cell.fill = styles["header"][1]
    row += 1
    recorded = list(export.checks.itertuples(index=False)) if not export.checks.empty else []
    for entry in recorded:
        sheet.cell(row, 1, _clean(entry.Check))
        sheet.cell(row, 2, _clean(entry.Detail)).alignment = styles["wrap"]
        sheet.cell(row, 3, _clean(entry.Value))
        row += 1
    for warning in warnings:
        sheet.cell(row, 1, "Workbook size")
        sheet.cell(row, 2, warning).alignment = styles["wrap"]
        row += 1
    if not recorded and not warnings:
        sheet.cell(row, 1, "No warnings were raised.")
    return sheet


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def build_workbook_bytes(
    export: ResourceFuelExport,
    *,
    pathways: PathwayExport | None = None,
    evidence: WorkbookEvidence | None = None,
) -> bytes:
    """Serialize the whole calculation as one traceable, formula-driven workbook."""
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - dependency guarded in requirements
        raise RuntimeError("XLSX export requires openpyxl.") from exc

    layout.validate_sheet_names()
    pathways = pathways or PathwayExport()
    evidence = evidence or WorkbookEvidence()
    warnings: list[str] = []

    workbook = Workbook()
    workbook.remove(workbook.active)

    _write_read_me(workbook, export)
    _write_glossary(workbook)
    _write_pooling_rules(workbook, evidence.pooling_rules)
    worked_example = _add_worked_example(workbook)

    # A run without flow evidence still needs the provenance columns to exist,
    # so the formulas on the sheets that point at them remain resolvable.
    flow_provenance = evidence.flow_provenance
    if flow_provenance is None or flow_provenance.empty:
        flow_provenance = pd.DataFrame(columns=list(_FLOW_PROVENANCE_LEAD))
    provenance = _write_value_table(
        workbook, layout.FLOW_PROVENANCE_SHEET,
        "Every reported flow, as first read from the run",
        "This is where the calculation starts. Each row is one connection in one model year, "
        "showing the value as reported, how many of the run's result rows were combined to reach "
        "it, and whether the connection exists in the run's own model connections.",
        flow_provenance, lead=_FLOW_PROVENANCE_LEAD, warnings=warnings,
    )
    _write_unit_conversion(workbook, evidence.unit_conversion, provenance)

    source_assignments = evidence.source_assignments
    if source_assignments is None or source_assignments.empty:
        source_assignments = pd.DataFrame(columns=list(_SOURCE_ASSIGNMENT_LEAD))
    assignment_anchor = _write_value_table(
        workbook, layout.SOURCE_COST_ASSIGNMENT_SHEET,
        "Every source-cost row and its resource-pool assignment",
        "One row is one reported cost assigned to one resource pool. A multi-output process can "
        "therefore appear more than once: Assignment share records the split and Assigned cost "
        "is the portion included in the pool total on sheet 11. Unallocated rows keep their reason.",
        source_assignments, lead=_SOURCE_ASSIGNMENT_LEAD, warnings=warnings,
    )

    raw_fuel = _write_value_table(
        workbook, layout.RAW_FUEL_USE_SHEET,
        "Fuel use, after time slices, model connections and units are settled",
        "One row per reported flow, once the earlier sheets have combined its result rows, "
        "checked it against the model connections, grouped its fuel and converted it to TWh. "
        "Rows marked as not included are shown so nothing is hidden, but the milestone sheet "
        "only adds up the included ones.",
        export.raw_fuel_use, lead=_RAW_FUEL_LEAD, warnings=warnings,
    )
    raw_source = _write_source_cost_summary(
        workbook,
        export.raw_source_costs,
        assignment_anchor,
        has_assignment_rows=not source_assignments.empty,
    )
    _write_value_table(
        workbook, layout.RAW_POOL_SUMMARY_SHEET,
        "The run's own resource pool summary",
        "Reproduced unchanged so the formulas in this workbook can be compared against the "
        "application's own figures.",
        evidence.pool_summary, warnings=warnings,
    )
    _write_value_table(
        workbook, layout.FUEL_TYPE_MAPPING_SHEET,
        "How each commodity was grouped",
        "The reported fuel type on the raw fuel use sheet follows from this grouping. Biomethane "
        "is separated from other bio fuels by name, then recombined for costing.",
        evidence.commodity_mapping, warnings=warnings,
    )
    _write_value_table(
        workbook, layout.MODEL_CONNECTIONS_SHEET,
        "Model connections between processes and commodities",
        "Used to establish which processes form the supply chain for each fuel.",
        evidence.model_connections, warnings=warnings,
    )
    raw_power = _write_value_table(
        workbook, layout.RAW_POWER_COST_SHEET,
        "Core cost rows on validated power-system processes",
        "These rows add up to the power-system core cost used to derive the electricity base cost.",
        export.raw_power_costs, lead=_RAW_POWER_LEAD, warnings=warnings,
    )
    _write_value_table(
        workbook, layout.COMMODITY_COST_SHEET,
        "Commodity costs, disclosed separately",
        "Cost_Com is reported here for completeness and is deliberately excluded from every "
        "formula in this workbook.",
        evidence.commodity_costs, warnings=warnings,
    )
    _write_value_table(
        workbook, layout.VINTAGE_DETAIL_SHEET,
        "Reported construction vintages",
        "Investment costs are split by the sector mix of the year the capacity was built. "
        "Fixed operating costs keep their reported vintage but are split by the reporting year's mix.",
        evidence.vintage_allocations, warnings=warnings,
    )
    _write_value_table(
        workbook, layout.RAW_ALLOCATION_SHEET,
        "The application's own allocation rows",
        "Kept so the live formulas here can be checked against the figures shown on screen.",
        evidence.allocations, warnings=warnings,
    )
    _write_value_table(
        workbook, layout.COMMODITY_COVERAGE_SHEET,
        "Which fuels count towards the cost split",
        "Every fuel that carries energy in this run, and whether its use helps decide how a "
        "pool's cost is divided. Anything marked as needing attention is not yet accounted for.",
        evidence.commodity_coverage, lead=_COVERAGE_LEAD, warnings=warnings,
    )

    _, milestone_fuel_column, milestone_fuel_rows = _write_milestone_fuel_use(
        workbook, export, raw_fuel
    )
    _, milestone_cost_column, milestone_cost_rows = _write_milestone_source_costs(
        workbook, export, raw_source
    )
    milestone_columns = {
        year: milestone_fuel_column + index for index, year in enumerate(export.milestone_years)
    }
    cost_milestone_columns = {
        year: milestone_cost_column + index for index, year in enumerate(export.milestone_years)
    }

    _, use_column, use_rows = _write_annual_fuel_use(
        workbook, export, milestone_columns, milestone_fuel_rows
    )
    _, cost_column, cost_rows = _write_annual_source_costs(
        workbook, export, cost_milestone_columns, milestone_cost_rows
    )
    _, share_column, share_rows = _write_allocation_shares(workbook, export, use_column, use_rows)
    _, allocated_column, allocated_rows, allocated_blocks = _write_allocated_costs(
        workbook, export, share_column, share_rows, cost_column, cost_rows
    )
    _, electricity_column, electricity_rows, electricity_sectors, electricity_block = _write_electricity(
        workbook, export, raw_power, use_column, use_rows, allocated_column, allocated_rows
    )
    _, resource_column, resource_rows = _write_resource_costs(
        workbook, export, allocated_column, allocated_rows, electricity_column, electricity_sectors
    )
    _write_fuelex(workbook, export, resource_column, resource_rows)

    _fill_worked_example(
        worked_example, export,
        provenance=provenance,
        use_column=use_column, use_rows=use_rows,
        cost_column=cost_column, cost_rows=cost_rows,
        share_column=share_column, share_rows=share_rows,
        allocated_column=allocated_column, allocated_rows=allocated_rows,
    )

    for name, title, subtitle, frame in (
        (layout.PATHWAY_LINKS_SHEET, "Physical pathway links",
         "Net supply to net demand for every model year, resource pool and individual fuel.",
         pathways.links),
        (layout.PATHWAY_PROCESS_NET_SHEET, "Process net calculation",
         "The net physical position of each process behind the pathway links.", pathways.process_net),
        (layout.PATHWAY_SOURCE_SHEET, "Source and output calculation",
         "Output-side evidence behind the supply end of each pathway.", pathways.source_detail),
        (layout.PATHWAY_DEMAND_SHEET, "Demand and input calculation",
         "Input-side evidence behind the demand end of each pathway.", pathways.demand_detail),
        (layout.PATHWAY_BALANCE_SHEET, "Exact commodity code balance audit",
         "Balances every exact commodity code within a pathway. No balancing flow is invented.",
         pathways.commodity_balance),
        (layout.PATHWAY_UNMATCHED_SHEET, "Unmatched reported rows",
         "Reported rows with no matching model connection. They are excluded from the graph.",
         pathways.unmatched),
        (layout.PATHWAY_HIDDEN_SHEET, "Rows hidden from the on-screen chart",
         "Valid physical rows kept out of the stakeholder chart for readability, retained in full here.",
         pathways.hidden),
        (layout.PATHWAY_NOTES_SHEET, "Physical pathway notes",
         "Technical notes recorded while building the physical pathways.", pathways.diagnostics),
        (layout.SUPPLY_CHAIN_LINKS_SHEET, "Supply-chain cost links",
         "Monetary supply-chain links behind the pool-level cost pathway view, for every model year.",
         pathways.supply_chain_links),
    ):
        _write_value_table(workbook, name, title, subtitle, frame, warnings=warnings)

    _write_checks(
        workbook, export, cost_column, cost_rows, allocated_column, allocated_blocks,
        electricity_column, electricity_rows, electricity_block, warnings,
    )
    _write_value_table(
        workbook, layout.TECHNICAL_NOTES_SHEET,
        "Technical allocation notes",
        "Notes recorded by the application while allocating costs in the selected run.",
        evidence.diagnostics, warnings=warnings,
    )

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_run_workbook(
    result,
    bundle,
    *,
    power_processes: tuple[str, ...] = (),
    reporter: object = None,
) -> bytes:
    """Build the whole workbook for one run, reporting progress as it goes.

    This runs on a worker thread, so it must never touch Streamlit.  It only
    reads the prepared allocation result and the run bundle.
    """
    from cost_allocation.resource_export import (
        prepare_all_fuel_pathways,
        prepare_flow_provenance,
        prepare_pooling_table,
        prepare_resource_fuel_export,
        prepare_unit_conversion,
    )
    from cost_allocation.resource_allocation import (
        load_resource_pools,
        summarise_commodity_coverage,
    )

    def stage(name: str, done: int = 0, total: int = 0) -> None:
        if reporter is not None:
            reporter.set(name, done, total)

    stage("Recalculating the allocation")
    export = prepare_resource_fuel_export(
        result,
        run_id=bundle.run_id,
        costs=bundle.costs,
        power_processes=power_processes,
        price_year=getattr(bundle, "price_year", None),
    )

    stage("Tracing fuel pathways")
    annual = getattr(bundle, "annual_flows", None)
    pathways = prepare_all_fuel_pathways(
        result,
        unmatched_flows=getattr(annual, "unmatched_edges", None),
        annual_diagnostics=getattr(annual, "diagnostics", None),
        progress=(reporter.advance if reporter is not None else None),
    )

    stage("Writing the workbook")
    evidence = WorkbookEvidence(
        pool_summary=result.pool_summary,
        commodity_mapping=result.commodity_mapping,
        model_connections=getattr(bundle, "vdt_topology", pd.DataFrame()),
        commodity_costs=result.commodity_cost_summary,
        vintage_allocations=result.vintage_allocations,
        allocations=result.allocations,
        diagnostics=result.diagnostics,
        flow_provenance=prepare_flow_provenance(annual),
        unit_conversion=prepare_unit_conversion(result.quantity_lineage),
        pooling_rules=prepare_pooling_table(load_resource_pools()),
        commodity_coverage=summarise_commodity_coverage(result.quantity_lineage),
        source_assignments=result.cost_assignments,
    )
    data = build_workbook_bytes(export, pathways=pathways, evidence=evidence)
    stage("Finished")
    return data

